"""Workbook profiling, behind a neutral backend interface.

The deterministic core profiles delimited text only. A real preclinical dataset
keeps its numbers in spreadsheets, so a core that cannot open a workbook reports
``tables: 0`` on a dataset full of tables -- and every downstream check that
depends on cell-level facts abstains for the wrong reason. That is D2A-47, found
by the XP14 run (#16).

Four commitments shape this module.

**The shape is observed, not interpreted.** A sheet's used range is a rectangle
the parser reports; this module records it and profiles what is inside. It does
not decide which sheet is "the data", does not merge multi-row headers, and does
not infer meaning from a column name.

**The header is decided by a named rule, and recorded as a decision.** A sheet
often has titles, merged banners or blank rows above the real header -- XP14's
own registry starts on row 2. The header row is chosen by
:mod:`data2agent.ingest.layout`, the same rule the delimited profiler uses
(D2A-97), or by a declaration when one names the sheet. The row used, the rows
skipped and why, and whether the choice was detected or declared are reported
in the profile, so the assumption is auditable rather than invisible.

**Missingness is resolved by the same machinery as CSV.** ``_observe`` from the
delimited profiler is reused deliberately, so a token means the same thing in a
sheet as in a CSV and ``manifest.missing_value_convention`` governs both. Two
implementations would let the convention drift, and the convention is the thing
that makes a missingness count reproducible.

**The backend is chosen by the bytes, and recorded.** OOXML (``xlsx``) is read
with openpyxl, as it always has been; genuine legacy BIFF ``xls``, ``xlsb`` and
``ods`` are read with python-calamine (D2A-94). Which format a file is was
decided by :mod:`data2agent.ingest.formats` from its content, never its name.
Every backend yields cells as the same Python types -- blank as ``None``, a
whole number as ``int``, a date as a midnight ``datetime`` -- so a sheet Excel
saved in four formats profiles identically in all four. Both backends read a
formula's *cached* value, not its expression; ``reader.cell_values`` records
that, so it is a stated limitation rather than a silent one.
"""

from __future__ import annotations

from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime, time
from importlib import metadata as _metadata
from pathlib import Path
from typing import Any

from ..errors import LayoutError
from ..ingest import layout as layouts
from ..ingest.conventions import DEFAULT_CONVENTION, MissingValueConvention
from ..ingest.tabular import (
    ColumnProfile,
    _convention_warnings,
    _observe,
    finalise,
    header_fields,
    new_column,
)

# Cap on sheets profiled per workbook. A pathological file should slow an
# ingest, not hang it; the cap is reported rather than applied silently.
_MAX_SHEETS = 64


class WorkbookReaderUnavailable(RuntimeError):
    """Raised when the optional parser is not installed."""


@dataclass(frozen=True)
class Backend:
    """One optional workbook parser and the formats it is trusted with."""

    name: str
    module: str  # import name
    distribution: str  # installed distribution, for the recorded version
    extra: str  # the data2agent extra that installs it
    formats: frozenset[str]

    def available(self) -> bool:
        try:
            __import__(self.module)
        except ImportError:
            return False
        return True

    def version(self) -> str | None:
        try:
            return _metadata.version(self.distribution)
        except _metadata.PackageNotFoundError:
            return None

    def describe(self) -> dict[str, str]:
        # Both backends open workbooks on cached values; neither exposes a
        # formula's expression. Recorded per sheet so a reader of the manifest
        # knows "43" typed and "43" computed are not distinguished here.
        return {"backend": self.name, "cell_values": "cached"}


OPENPYXL = Backend("openpyxl", "openpyxl", "openpyxl", "xlsx", frozenset({"xlsx"}))
CALAMINE = Backend(
    "calamine",
    "python_calamine",
    "python-calamine",
    "workbooks",
    frozenset({"xls", "xlsb", "ods"}),
)
BACKENDS: dict[str, Backend] = {OPENPYXL.name: OPENPYXL, CALAMINE.name: CALAMINE}


def backend_for(format_id: str) -> Backend:
    """The backend that reads ``format_id``; ``KeyError`` for any other format.

    OOXML stays on openpyxl even when calamine is installed, so an existing
    XLSX manifest does not change because an unrelated extra was added.
    """
    for backend in BACKENDS.values():
        if format_id in backend.formats:
            return backend
    raise KeyError(f"no workbook backend reads format {format_id!r}")


@dataclass
class SheetProfile:
    """Observed properties of one worksheet."""

    path: str  # "<workbook path>#<sheet name>"
    workbook: str
    sheet: str
    sheet_index: int
    sheet_state: str  # "visible" | "hidden" | "veryHidden"
    header_row: int | None
    has_header: bool
    # None, never 0, when the content was not observed. Zero rows is a finding
    # about a sheet that was read; an unreadable file has no row count at all,
    # and reporting one turns an inability to profile into false evidence.
    profiled: bool
    rows: int | None
    columns: list[ColumnProfile]
    merged_ranges: int
    convention: MissingValueConvention
    warnings: list[str] = field(default_factory=list)
    reader: Backend = OPENPYXL
    layout: layouts.HeaderLayout | None = None

    def as_dict(self) -> dict[str, object]:
        return {
            "path": self.path,
            "workbook": self.workbook,
            "sheet": self.sheet,
            "sheet_index": self.sheet_index,
            "sheet_state": self.sheet_state,
            **header_fields(self.layout),
            "has_header": self.has_header,
            "profiled": self.profiled,
            "rows": self.rows,
            "column_count": len(self.columns),
            "merged_ranges": self.merged_ranges,
            "reader": self.reader.describe(),
            "missing_convention": self.convention.as_dict(),
            "columns": [column.as_dict() for column in self.columns],
            "missing": {column.name: column.missing for column in self.columns},
            "warnings": list(self.warnings),
        }


def available(format_id: str = "xlsx") -> bool:
    """Whether the optional parser for ``format_id`` is importable."""
    return backend_for(format_id).available()


def profile_workbook(
    path: Path,
    relative_path: str,
    convention: MissingValueConvention = DEFAULT_CONVENTION,
    format_id: str = "xlsx",
    layout_for: Callable[[str], layouts.TableLayout | None] | None = None,
) -> list[SheetProfile]:
    """Profile every worksheet in a workbook whose format the bytes established.

    Raises :class:`WorkbookReaderUnavailable` when the optional parser is
    missing, so the caller can report a workbook it could not read instead of
    reporting a dataset with no tables. ``layout_for`` maps a sheet's table path
    to its declared layout, if any; a declaration the sheet cannot satisfy
    raises :class:`LayoutError` rather than degrading to an unreadable sheet.
    """
    backend = backend_for(format_id)
    if not backend.available():
        raise WorkbookReaderUnavailable(_unavailable_message(backend))
    try:
        book = _open(path, backend)
    except Exception as exc:  # parsers raise a wide range on malformed input
        return [_unreadable(relative_path, convention, f"{type(exc).__name__}: {exc}", backend)]

    profiles: list[SheetProfile] = []
    try:
        sheet_names = book.sheet_names()
        for index, name in enumerate(sheet_names):
            if index >= _MAX_SHEETS:
                profiles.append(
                    _unreadable(
                        relative_path,
                        convention,
                        f"workbook has {len(sheet_names)} sheets; "
                        f"only the first {_MAX_SHEETS} were profiled",
                        backend,
                    )
                )
                break
            declared = layout_for(f"{relative_path}#{name}") if layout_for else None
            try:
                profiles.append(
                    _profile_sheet(book, name, relative_path, index, convention, declared)
                )
            except LayoutError:
                raise  # a wrong declaration is the user's to fix, not a bad sheet
            except Exception as exc:
                # read_only defers XML parsing until the rows are iterated, so a
                # sheet can fail long after the workbook opened cleanly. One bad
                # sheet must not abort the ingest of the whole dataset.
                profiles.append(
                    _unreadable(
                        f"{relative_path}#{name}",
                        convention,
                        f"sheet '{name}' could not be read: {type(exc).__name__}: {exc}",
                        backend,
                        workbook=relative_path,
                        sheet=name,
                    )
                )
    finally:
        book.close()

    return profiles


class OpenWorkbook:
    """The neutral surface profiling and row access use, whatever the backend."""

    backend: Backend

    def close(self) -> None:
        raise NotImplementedError

    def sheet_names(self) -> list[str]:
        raise NotImplementedError

    def sheet_state(self, name: str) -> str:
        raise NotImplementedError

    def merged_ranges(self, name: str) -> int:
        raise NotImplementedError

    def iter_rows(self, name: str, min_row: int = 1) -> Iterator[tuple[Any, ...]]:
        """Rows from ``min_row``, numbered as the spreadsheet numbers them (1-based)."""
        raise NotImplementedError


@contextmanager
def open_workbook(path: Path, backend: Backend) -> Iterator[OpenWorkbook]:
    """Open ``path`` read-only with ``backend``. The file is never converted or written."""
    if not backend.available():
        raise WorkbookReaderUnavailable(_unavailable_message(backend))
    book = _open(path, backend)
    try:
        yield book
    finally:
        book.close()


def _open(path: Path, backend: Backend) -> OpenWorkbook:
    return _OpenpyxlWorkbook(path) if backend is OPENPYXL else _CalamineWorkbook(path)


def _unavailable_message(backend: Backend) -> str:
    return (
        f"reading this workbook requires the '{backend.extra}' extra: "
        f"pip install 'data2agent[{backend.extra}]'"
    )


class _OpenpyxlWorkbook(OpenWorkbook):
    backend = OPENPYXL

    def __init__(self, path: Path) -> None:
        import openpyxl

        # Opened from a file handle, not a path: openpyxl dispatches on the
        # extension and refuses an OOXML workbook named '.xls', which is the case
        # D2A-46 identified in XP14. A handle is both extension-independent and
        # seekable, so read_only streaming still applies -- reading the archive
        # into memory first would defeat it on exactly the large files it protects.
        self._handle = path.open("rb")
        try:
            self._book = openpyxl.load_workbook(
                self._handle,
                read_only=True,  # streaming: a large workbook must not be held whole
                data_only=True,  # cached values, not formula text; see module note
            )
        except Exception:
            self._handle.close()
            raise

    def close(self) -> None:
        self._book.close()
        self._handle.close()

    def sheet_names(self) -> list[str]:
        return list(self._book.sheetnames)

    def _sheet(self, name: str):
        if name not in self._book.sheetnames:
            raise KeyError(f"worksheet {name!r} no longer exists in the workbook")
        return self._book[name]

    def sheet_state(self, name: str) -> str:
        return getattr(self._sheet(name), "sheet_state", "visible") or "visible"

    def merged_ranges(self, name: str) -> int:
        # read_only worksheets expose merged ranges inconsistently; absence is
        # not evidence of none, so it is reported as 0 without claiming certainty.
        merged = getattr(self._sheet(name), "merged_cells", None)
        return len(merged.ranges) if merged is not None and hasattr(merged, "ranges") else 0

    def iter_rows(self, name: str, min_row: int = 1) -> Iterator[tuple[Any, ...]]:
        return self._sheet(name).iter_rows(min_row=min_row, values_only=True)


class _CalamineWorkbook(OpenWorkbook):
    backend = CALAMINE

    def __init__(self, path: Path) -> None:
        from python_calamine import CalamineWorkbook

        # calamine has no write API at all, so the source cannot be converted or
        # touched in place; it is opened by content, never by extension.
        with _panics_as_errors():
            self._book = CalamineWorkbook.from_path(str(path))
            self._states = {
                meta.name: _CALAMINE_STATES.get(str(meta.visible).rsplit(".", 1)[-1], "unknown")
                for meta in self._book.sheets_metadata
            }

    def close(self) -> None:
        self._book.close()

    def sheet_names(self) -> list[str]:
        return list(self._book.sheet_names)

    def _sheet(self, name: str):
        if name not in self._book.sheet_names:
            raise KeyError(f"worksheet {name!r} no longer exists in the workbook")
        with _panics_as_errors():
            return self._book.get_sheet_by_name(name)

    def sheet_state(self, name: str) -> str:
        return self._states.get(name, "unknown")

    def merged_ranges(self, name: str) -> int:
        # None where the format reader does not track merges (xlsb, ods):
        # unknown, and reported as 0 exactly as openpyxl's read_only mode is.
        merged = self._sheet(name).merged_cell_ranges
        return len(merged) if merged else 0

    def iter_rows(self, name: str, min_row: int = 1) -> Iterator[tuple[Any, ...]]:
        # calamine pads leading empty *rows*, so the row number here is the
        # spreadsheet's own; it does not pad leading empty *columns*, so each row
        # is left-padded to column A. Without that, a sheet starting at column C
        # would shift every column position relative to openpyxl.
        sheet = self._sheet(name)
        if sheet.start is None:
            # An empty sheet has no used range, and calamine's iter_rows panics
            # on it (a Rust unwrap) instead of yielding nothing. Found on the
            # blank 'Sheet1' every XP1 ECG .xlsb export carries.
            return
        lead = (None,) * sheet.start[1]
        with _panics_as_errors():
            for number, row in enumerate(sheet.iter_rows(), start=1):
                if number >= min_row:
                    yield lead + tuple(_calamine_cell(value) for value in row)


@contextmanager
def _panics_as_errors() -> Iterator[None]:
    """Turn a Rust panic inside calamine into an ordinary, catchable error.

    pyo3 raises ``PanicException``, a ``BaseException``: it escapes every
    ``except Exception`` and would abort a whole ingest on one malformed sheet.
    It is matched by name because pyo3 does not export it as an importable class.
    """
    try:
        yield
    except BaseException as exc:
        if type(exc).__name__ != "PanicException":
            raise
        raise RuntimeError(f"calamine could not parse this content: {exc}") from exc


_CALAMINE_STATES = {"Visible": "visible", "Hidden": "hidden", "VeryHidden": "veryHidden"}


def _calamine_cell(value: Any) -> Any:
    """Bring a calamine cell to the type openpyxl yields for the same cell.

    A workbook stores every number as a double and every date as a serial day,
    so these are not reinterpretations: an integral double is the integer the
    sheet displays, and a date is that day's midnight.
    """
    if isinstance(value, str) and value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, time())
    return value


def _unreadable(
    relative_path: str,
    convention: MissingValueConvention,
    note: str,
    backend: Backend = OPENPYXL,
    *,
    workbook: str | None = None,
    sheet: str = "",
) -> SheetProfile:
    """A placeholder recording that a workbook exists and was not profiled.

    Emitted rather than returning nothing, because an absent table and an
    unreadable one are different facts and only one of them is a data quality
    finding about the dataset.
    """
    return SheetProfile(
        path=relative_path,
        workbook=workbook if workbook is not None else relative_path,
        sheet=sheet,
        sheet_index=-1,
        sheet_state="unknown",
        header_row=None,
        has_header=False,
        profiled=False,
        rows=None,
        columns=[],
        merged_ranges=0,
        convention=convention,
        warnings=[note],
        reader=backend,
    )


def _profile_sheet(
    book: OpenWorkbook,
    name: str,
    workbook_path: str,
    index: int,
    convention: MissingValueConvention,
    declared: layouts.TableLayout | None = None,
) -> SheetProfile:
    warnings: list[str] = []
    state = book.sheet_state(name)
    if state != "visible":
        warnings.append(
            f"sheet is '{state}'; profiled anyway, because a hidden sheet is "
            f"still data the file carries"
        )

    merged = book.merged_ranges(name)

    columns: list[ColumnProfile] = []
    data_rows = 0

    numbered = (
        layouts.Row(number, raw_row) for number, raw_row in enumerate(book.iter_rows(name), start=1)
    )
    layout, body = layouts.split_header(numbered, declared)
    warnings.extend(layout.warnings)
    header_row_index = layout.header_row
    if header_row_index is not None:
        header = _header_names(layout.labels, warnings, name, header_row_index)
        columns = [new_column(n, i) for i, n in enumerate(header)]
        if layout.header_cells is not None:
            for column, cells in zip(columns, layout.header_cells, strict=False):
                column.header_cells = cells

    # Every row from the data start is an observation, blank ones included, as
    # before D2A-97: the row reader counts offsets over the same rows.
    for row in body:
        cells = ["" if v is None else str(v) for v in row.cells]
        data_rows += 1
        if len(cells) > len(columns):
            # Cells to the right of the header row. Recorded, never discarded:
            # a header narrower than its data is a finding about the sheet.
            for extra in range(len(columns), len(cells)):
                columns.append(new_column(_column_letter(extra), extra))
                columns[-1].missing_empty += data_rows - 1
        for position, column in enumerate(columns):
            _observe(column, cells[position] if position < len(cells) else "", convention)

    # Finalised by the delimited profiler's own routine, so the invariant
    # missing == missing_empty + missing_sentinel, and the uniqueness verdict a
    # metadata rule reads, are computed identically for a sheet and for a CSV.
    for column in columns:
        finalise(column, data_rows)

    warnings.extend(_convention_warnings(columns, convention))

    if header_row_index is None:
        warnings.append("sheet is empty; no header row and no data rows")

    return SheetProfile(
        path=f"{workbook_path}#{name}",
        workbook=workbook_path,
        sheet=name,
        sheet_index=index,
        sheet_state=state,
        header_row=header_row_index,
        has_header=header_row_index is not None,
        profiled=True,
        rows=data_rows,
        columns=columns,
        merged_ranges=merged,
        convention=convention,
        warnings=warnings,
        reader=book.backend,
        layout=layout,
    )


def _header_names(cells: list[str], warnings: list[str], sheet: str, row: int) -> list[str]:
    """Name the columns from the header row, filling blanks positionally.

    A blank header cell is named for its spreadsheet column so that the column
    is still profiled. Silently dropping it would under-report the sheet's width.
    """
    names: list[str] = []
    blanks = 0
    # Tracks what has been EMITTED, not what was read. Suffixing against the
    # original labels lets ['id', 'id', 'id.1'] collapse to two identical
    # 'id.1' entries, which then collide in the name-keyed missing map.
    emitted: set[str] = set()
    for position, cell in enumerate(cells):
        label = cell.strip()
        if not label:
            blanks += 1
            label = _column_letter(position)
        if label in emitted:
            base, suffix = label, 1
            while f"{base}.{suffix}" in emitted:
                suffix += 1
            label = f"{base}.{suffix}"
        emitted.add(label)
        names.append(label)

    while names and names[-1] == _column_letter(len(names) - 1):
        names.pop()  # trailing empties are the edge of the used range, not columns

    if blanks:
        warnings.append(
            f"header row {row} has {blanks} blank cell(s); those columns are named "
            f"for their spreadsheet column letter. The row used, and how it was "
            f"chosen, are recorded in 'header_row' and 'header_source'; declare the "
            f"layout (--layout) if it is not the real header"
        )
    return names


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters
