"""Shared, data-free machinery for the XP16 data-plane benchmark.

Nothing in this module names a file, sheet, column, identifier or value of the
dataset. Everything dataset-specific lives in the package's local, git-ignored
``config/tables.yaml`` and ``config/questions.spec.yaml``; this module only
knows how to read tables *as declared there* and how to evaluate a small, closed
set of gold computations over them.

Independence from the system under test is the point. Tables are read with
``openpyxl`` and the standard-library ``csv`` module, never through
``data2agent``, so the tool is not graded against itself. The only coupling is
the optional :class:`RowSource` protocol, which lets the baseline script feed
the *same* evaluator with rows obtained from the Data2Agent service instead --
the computation is then identical and only the retrieval differs, which is
exactly the thing being measured.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import os
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Protocol


def _repo_root() -> Path:
    env = os.environ.get("D2A_REPO")
    return Path(env).resolve() if env else Path(__file__).resolve().parents[3]


REPO = _repo_root()

# Rounding applied to every computed float before it is written to the gold. It
# is far below any tolerance the questions use and exists only so that re-runs
# are byte-identical across platforms (a float's repr can differ in its last ulp
# depending on summation order, and the gold must not).
FLOAT_DIGITS = 6

ABSTAIN = "ABSTAIN"
PENDING_PREFIX = "PENDING-"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1 << 20):
            digest.update(chunk)
    return digest.hexdigest()


def load_yaml(path: Path) -> Any:
    import yaml  # noqa: PLC0415 -- optional dependency, only the gold chain needs it

    with path.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def dump_yaml(payload: Any) -> str:
    import yaml  # noqa: PLC0415

    return yaml.safe_dump(payload, sort_keys=False, allow_unicode=True, width=100)


# --------------------------------------------------------------------------
# value normalisation
# --------------------------------------------------------------------------


def norm_value(value: Any) -> Any:
    """Canonical JSON-safe form of a cell value, identical on every run."""
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isfinite(value) and value == int(value) and abs(value) < 1e15:
            return int(value)
        return round(value, FLOAT_DIGITS)
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value
    return str(value)


def to_number(value: Any) -> float | None:
    """A cell as a number, or None when it is not one. Nothing is coerced silently:
    a text marker such as '-' or '?' stays non-numeric and is reported by callers."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def is_blank(value: Any) -> bool:
    """None, or a string holding only whitespace. Never a unit, a group or a value."""
    return value is None or (isinstance(value, str) and not value.strip())


def round_num(x: float) -> int | float:
    r = round(float(x), FLOAT_DIGITS)
    return int(r) if r == int(r) and abs(r) < 1e15 else r


# --------------------------------------------------------------------------
# tables
# --------------------------------------------------------------------------


@dataclass
class Row:
    values: dict[str, Any]
    locator: int  # 1-based worksheet row, or 1-based physical line for delimited files
    cells: dict[str, str] = field(default_factory=dict)  # column name -> "B4" / "L4:F3"


@dataclass
class Table:
    table_id: str
    file: str  # path relative to the package's source/
    sha256: str
    sheet: str | None
    columns: list[str]
    column_letters: dict[str, str]  # column name -> spreadsheet letter / 1-based field index
    rows: list[Row]
    header_rows: list[int]
    # 'row' (a sheet row / a physical line) or 'interval' (the n-th behaviour
    # interval a BORIS project yields, see FileRowSource._read_boris)
    locator_kind: str = "row"

    def locator_range(self, column: str | None = None, rows: Iterable[Row] | None = None) -> str:
        chosen = list(rows) if rows is not None else self.rows
        if not chosen:
            return ""
        lo = min(r.locator for r in chosen)
        hi = max(r.locator for r in chosen)
        if self.locator_kind == "interval":
            return f"intervals {lo}-{hi}" + (f", column {column}" if column else "")
        if self.sheet is not None:
            if column is None:
                return f"{lo}:{hi}"
            letter = self.column_letters[column]
            return f"{letter}{lo}:{letter}{hi}"
        if column is None:
            return f"lines {lo}-{hi}"
        return f"lines {lo}-{hi}, field {self.column_letters[column]}"


class RowSource(Protocol):
    """Where the evaluator gets a declared table from."""

    def table(self, table_id: str) -> Table: ...

    def raw_cell(self, file: str, sheet: str, cell: str) -> Any: ...

    def file_sha(self, file: str) -> str: ...


def _col_letter(index: int) -> str:
    letters = ""
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for name in names:
        if name in seen:
            seen[name] += 1
            out.append(f"{name} #{seen[name]}")
        else:
            seen[name] = 1
            out.append(name)
    return out


def _compose_headers(
    header_matrix: list[list[Any]], *, ffill: bool, drop: set[str]
) -> list[str | None]:
    """Combine one or more header rows into one name per column.

    Upper rows may be forward-filled (a banner spanning several columns). Values
    listed in ``drop`` (preamble labels) never contribute to a name. A column
    whose every header cell is empty has no name and is not exposed.
    """
    width = max((len(r) for r in header_matrix), default=0)
    parts_per_col: list[list[str]] = [[] for _ in range(width)]
    for depth, raw in enumerate(header_matrix):
        is_last = depth == len(header_matrix) - 1
        carry: str | None = None
        for i in range(width):
            v = raw[i] if i < len(raw) else None
            text = "" if v is None else str(norm_value(v)).strip()
            if text in drop:
                text = ""
            if text:
                carry = text
            elif ffill and not is_last and carry is not None:
                text = carry
            if text:
                parts_per_col[i].append(text)
    return [" | ".join(p) if p else None for p in parts_per_col]


class FileRowSource:
    """Reads declared tables straight from the frozen bytes (openpyxl / csv)."""

    def __init__(self, source_dir: Path, config: dict[str, Any]) -> None:
        self.source_dir = source_dir
        self.config = config
        self._tables: dict[str, Table] = {}
        self._books: dict[str, Any] = {}
        self._shas: dict[str, str] = {}

    def file_sha(self, file: str) -> str:
        if file not in self._shas:
            self._shas[file] = sha256_file(self.source_dir / file)
        return self._shas[file]

    def _book(self, file: str) -> Any:
        if file not in self._books:
            import openpyxl  # noqa: PLC0415

            data = (self.source_dir / file).read_bytes()
            # by bytes, so the extension never decides how a file is opened
            self._books[file] = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        return self._books[file]

    def raw_cell(self, file: str, sheet: str, cell: str) -> Any:
        return norm_value(self._book(file)[sheet][cell].value)

    def sheet_names(self, file: str) -> list[str]:
        return list(self._book(file).sheetnames)

    def table(self, table_id: str) -> Table:
        if table_id in self._tables:
            return self._tables[table_id]
        spec = self.config["tables"].get(table_id)
        if spec is None:
            raise KeyError(f"table {table_id!r} is not declared in tables.yaml")
        kind = spec.get("kind", "xlsx")
        if kind == "xlsx":
            table = self._read_xlsx(table_id, spec)
        elif kind == "delimited":
            table = self._read_delimited(table_id, spec)
        elif kind == "boris":
            table = self._read_boris(table_id, spec)
        else:
            raise ValueError(f"table {table_id!r}: unknown kind {kind!r}")
        _apply_derived(table, spec.get("derived", {}))
        self._tables[table_id] = table
        return table

    def _read_xlsx(self, table_id: str, spec: dict[str, Any]) -> Table:
        ws = self._book(spec["file"])[spec["sheet"]]
        header_rows = list(spec["header_rows"])
        max_col = ws.max_column
        matrix = [[ws.cell(r, c).value for c in range(1, max_col + 1)] for r in header_rows]
        names = _compose_headers(
            matrix, ffill=spec.get("ffill_upper_headers", True), drop=set(spec.get("drop", []))
        )
        first = int(spec.get("first_row", header_rows[-1] + 1))
        last = spec.get("last_row")
        col_index = [(i + 1, n) for i, n in enumerate(names) if n is not None]
        unique = _dedupe([n for _, n in col_index])
        col_index = [(ci, un) for (ci, _), un in zip(col_index, unique, strict=True)]
        letters = {name: _col_letter(ci) for ci, name in col_index}
        rows: list[Row] = []
        r = first
        stop = int(last) if last is not None else ws.max_row
        while r <= stop:
            values = {name: norm_value(ws.cell(r, ci).value) for ci, name in col_index}
            if all(v is None or v == "" for v in values.values()):
                if last is None:
                    break
            else:
                rows.append(Row(values, r, {name: f"{letters[name]}{r}" for name in values}))
            r += 1
        return Table(
            table_id,
            spec["file"],
            self.file_sha(spec["file"]),
            spec["sheet"],
            [n for _, n in col_index],
            letters,
            rows,
            header_rows,
        )

    def _read_boris(self, table_id: str, spec: dict[str, Any]) -> Table:
        """One row per behaviour interval of a BORIS project (JSON).

        A BORIS state event toggles: the first occurrence of a (subject,
        behaviour) starts it, the next stops it. Occurrences are paired in time
        order per (subject, behaviour), exactly as that rule states; a point
        event is an interval of length 0. A start with no matching stop is kept
        as a row with an empty stop and duration, never closed by assumption.
        ``observations`` (a regex, optional) limits the observations read.
        """
        import json  # noqa: PLC0415

        doc = json.loads((self.source_dir / spec["file"]).read_text(encoding="utf-8"))
        types = {b.get("code"): b.get("type", "") for b in doc.get("behaviors_conf", {}).values()}
        pattern = re.compile(spec["observations"]) if spec.get("observations") else None
        columns = [
            "Observation id",
            "Subject",
            "Behavior",
            "Type",
            "Start (s)",
            "Stop (s)",
            "Duration (s)",
        ]
        rows: list[Row] = []
        for obs_id in sorted(doc.get("observations", {})):
            if pattern and not pattern.fullmatch(obs_id):
                continue
            events = doc["observations"][obs_id].get("events", [])
            ordered = sorted(enumerate(events), key=lambda ie: (float(ie[1][0]), ie[0]))
            open_: dict[tuple[str, str], float] = {}
            intervals: list[tuple[str, str, str, float, float | None]] = []
            for _, ev in ordered:
                t, subject, code = float(ev[0]), str(ev[1]), str(ev[2])
                if "state" not in types.get(code, "").lower():
                    intervals.append((subject, code, types.get(code, ""), t, t))
                elif (subject, code) in open_:
                    intervals.append((subject, code, types[code], open_.pop((subject, code)), t))
                else:
                    open_[(subject, code)] = t
            for (subject, code), t in open_.items():
                intervals.append((subject, code, types.get(code, ""), t, None))
            for subject, code, typ, start, stop in sorted(intervals, key=lambda x: (x[3], x[1])):
                n = len(rows) + 1
                values = {
                    "Observation id": obs_id,
                    "Subject": subject,
                    "Behavior": code,
                    "Type": typ,
                    "Start (s)": norm_value(start),
                    "Stop (s)": None if stop is None else norm_value(stop),
                    "Duration (s)": None if stop is None else round_num(stop - start),
                }
                rows.append(Row(values, n, {c: f"interval {n}" for c in columns}))
        return Table(
            table_id,
            spec["file"],
            self.file_sha(spec["file"]),
            None,
            columns,
            {c: c for c in columns},
            rows,
            [],
            locator_kind="interval",
        )

    def _read_delimited(self, table_id: str, spec: dict[str, Any]) -> Table:
        raw = (self.source_dir / spec["file"]).read_bytes()
        text = raw.decode(spec.get("encoding", "utf-8"))
        lines = list(csv.reader(io.StringIO(text, newline=""), delimiter=spec["delimiter"]))
        header_rows = list(spec["header_rows"])
        matrix = [lines[i - 1] for i in header_rows]
        names = _compose_headers(
            matrix, ffill=spec.get("ffill_upper_headers", False), drop=set(spec.get("drop", []))
        )
        col_index = [(i, n) for i, n in enumerate(names) if n is not None]
        unique = _dedupe([n for _, n in col_index])
        col_index = [(ci, un) for (ci, _), un in zip(col_index, unique, strict=True)]
        letters = {name: str(ci + 1) for ci, name in col_index}
        first = int(spec.get("first_row", header_rows[-1] + 1))
        last = spec.get("last_row")  # honoured exactly as for workbooks
        rows = []
        for lineno, record in enumerate(lines, start=1):
            if last is not None and lineno > int(last):
                break
            if lineno < first or not record or all(not f.strip() for f in record):
                continue
            values = {name: (record[ci] if ci < len(record) else None) for ci, name in col_index}
            rows.append(
                Row(
                    values,
                    lineno,
                    {name: f"line {lineno}, field {letters[name]}" for name in values},
                )
            )
        return Table(
            table_id,
            spec["file"],
            self.file_sha(spec["file"]),
            None,
            [n for _, n in col_index],
            letters,
            rows,
            header_rows,
        )


def _apply_derived(table: Table, derived: dict[str, str]) -> None:
    """Template columns such as ``"{Cage}-{Tail}"``; a missing part yields None."""
    for name, template in derived.items():
        fields = re.findall(r"{([^}]+)}", template)
        for row in table.rows:
            parts = {f: row.values.get(f) for f in fields}
            if any(v is None or v == "" for v in parts.values()):
                row.values[name] = None
            else:
                out = template
                for f, v in parts.items():
                    out = out.replace("{" + f + "}", str(v))
                row.values[name] = out
            row.cells[name] = "+".join(row.cells.get(f, "?") for f in fields)
        table.columns.append(name)
        table.column_letters[name] = "+".join(table.column_letters.get(f, "?") for f in fields)


# --------------------------------------------------------------------------
# identifier transforms (declared crosswalks)
# --------------------------------------------------------------------------


def make_transform(spec: dict[str, Any] | list[dict[str, Any]] | None) -> Callable[[Any], Any]:
    """A declared identifier rewrite: ``{regex, template}`` or an ordered list of
    them (first full match wins). A value no alternative matches maps to None,
    so an unexpected form surfaces as unmatched instead of passing through."""
    if not spec:
        return lambda v: v
    alternatives = spec if isinstance(spec, list) else [spec]
    compiled = [(re.compile(a["regex"]), a["template"]) for a in alternatives]

    def apply(value: Any) -> Any:
        if value is None:
            return None
        for pattern, template in compiled:
            m = pattern.fullmatch(str(norm_value(value)))
            if m:
                return template.format(**m.groupdict())
        return None

    return apply


# --------------------------------------------------------------------------
# filters
# --------------------------------------------------------------------------


def _match(row: Row, rule: dict[str, Any]) -> bool:
    v = row.values.get(rule["col"])
    op = rule.get("op", "eq")
    target = rule.get("value")
    if op == "eq":
        return norm_value(v) == target or (
            to_number(v) is not None and to_number(v) == to_number(target)
        )
    if op == "ne":
        return not _match(row, {**rule, "op": "eq"})
    if op == "in":
        return any(_match(row, {"col": rule["col"], "op": "eq", "value": t}) for t in target)
    if op == "not_in":
        return not _match(row, {**rule, "op": "in"})
    if op == "empty":
        return v is None or (isinstance(v, str) and not v.strip())
    if op == "nonempty":
        return not (v is None or (isinstance(v, str) and not v.strip()))
    if op == "numeric":
        return to_number(v) is not None
    if op == "not_numeric":
        return to_number(v) is None
    if op == "regex":
        return v is not None and re.fullmatch(target, str(v)) is not None
    if op in {"gt", "ge", "lt", "le"}:
        x, t = to_number(v), to_number(target)
        if x is None or t is None:
            return False
        return {"gt": x > t, "ge": x >= t, "lt": x < t, "le": x <= t}[op]
    if op in {"span_ge", "span_lt"}:
        # an interval cell 'start-end' (e.g. a time bin '300.000-600.000'),
        # compared by its length; a cell that is not such an interval never matches
        span, t = _span(v), to_number(target)
        if span is None or t is None:
            return False
        tol = 1e-6
        return span >= t - tol if op == "span_ge" else span < t - tol
    raise ValueError(f"unknown filter op {op!r}")


def _span(value: Any) -> float | None:
    m = re.fullmatch(r"\s*([0-9]+(?:\.[0-9]*)?)\s*-\s*([0-9]+(?:\.[0-9]*)?)\s*", str(value or ""))
    if not m:
        return None
    return float(m.group(2)) - float(m.group(1))


def select(table: Table, where: list[dict[str, Any]] | None) -> list[Row]:
    rules = where or []
    for rule in rules:
        if rule["col"] not in table.columns:
            raise KeyError(f"{table.table_id}: no column {rule['col']!r}; have {table.columns}")
    return [r for r in table.rows if all(_match(r, rule) for rule in rules)]


# --------------------------------------------------------------------------
# the gold computations
# --------------------------------------------------------------------------


class GoldError(Exception):
    """A computation that could not be carried out as declared."""


@dataclass
class Result:
    answer: Any
    sources: list[dict[str, Any]]
    computation: str


def _src(table: Table, rng: str, rows: list[Row] | None = None) -> dict[str, Any]:
    out: dict[str, Any] = {"file": table.file, "sha256": table.sha256}
    if table.sheet is not None:
        out["sheet"] = table.sheet
    out["range"] = rng
    if rows is not None:
        out["rows"] = sorted({r.locator for r in rows})
    return out


def _need(table: Table, *cols: str) -> None:
    missing = [c for c in cols if c not in table.columns]
    if missing:
        raise KeyError(f"{table.table_id}: no column(s) {missing}; have {table.columns}")


def _stats(values: list[float]) -> dict[str, Any]:
    n = len(values)
    mean = sum(values) / n if n else None
    if n > 1:
        var = sum((x - mean) ** 2 for x in values) / (n - 1)
        sd = math.sqrt(var)
        sem = sd / math.sqrt(n)
    else:
        sd = sem = None
    return {
        "n": n,
        "mean": round_num(mean) if mean is not None else None,
        "sd": round_num(sd) if sd is not None else None,
        "sem": round_num(sem) if sem is not None else None,
    }


def _attach(
    src: RowSource,
    config: dict[str, Any],
    table: Table,
    rows: list[Row],
    join: dict[str, Any],
    sources: list[dict[str, Any]],
) -> tuple[list[Row], list[str]]:
    """Left-join columns from another declared table onto ``rows``.

    Keys may pass through a declared transform on either side. Unmatched left
    rows are returned by unit so the caller can refuse or report them; nothing
    is dropped quietly.
    """
    other = src.table(join["table"])
    _need(table, join["left_key"])
    _need(other, join["right_key"], *join["columns"])
    lt = make_transform(config.get("transforms", {}).get(join.get("left_transform", "")))
    rt = make_transform(config.get("transforms", {}).get(join.get("right_transform", "")))
    index: dict[Any, Row] = {}
    for r in select(other, join.get("where")):
        k = rt(r.values.get(join["right_key"]))
        if k is None:
            continue
        if k in index:
            # A repeated right key is allowed only when declared (right_rows_agree)
            # AND every repeat carries the same joined values -- e.g. several time
            # bins of one observation that all name the same animal and genotype.
            # Anything else would be a choice between rows, and is refused.
            same = all(
                norm_value(index[k].values.get(c)) == norm_value(r.values.get(c))
                for c in join["columns"]
            )
            if not (join.get("right_rows_agree") and same):
                raise GoldError(
                    f"join key {k!r} is not unique in {other.table_id}"
                    + ("" if same else " and its rows disagree on the joined columns")
                )
            continue
        index[k] = r
    out: list[Row] = []
    unmatched: list[str] = []
    used: list[Row] = []
    for r in rows:
        k = lt(r.values.get(join["left_key"]))
        hit = index.get(k)
        if hit is None:
            unmatched.append(str(r.values.get(join["left_key"])))
            continue
        used.append(hit)
        merged = dict(r.values)
        cells = dict(r.cells)
        for c in join["columns"]:
            alias = join.get("as", {}).get(c, c)
            merged[alias] = hit.values.get(c)
            cells[alias] = hit.cells.get(c, "")
        out.append(Row(merged, r.locator, cells))
    for c in [join["right_key"], *join["columns"]]:
        sources.append(_src(other, other.locator_range(c, used), used))
    return out, unmatched


def evaluate(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    op = spec["op"]
    fn = _OPS.get(op)
    if fn is None:
        raise GoldError(f"unknown op {op!r}")
    return fn(src, config, spec)


def _op_cell(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    t = src.table(spec["table"])
    _need(t, spec["column"])
    rows = select(t, spec.get("where"))
    if len(rows) != 1:
        raise GoldError(f"cell: expected exactly one row in {t.table_id}, matched {len(rows)}")
    row = rows[0]
    value = row.values[spec["column"]]
    return Result(
        value,
        [_src(t, row.cells[spec["column"]], [row])],
        f"value of column '{spec['column']}' in the single row of table '{t.table_id}' "
        f"matching {spec.get('where')}",
    )


def _op_raw_cell(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    value = src.raw_cell(spec["file"], spec["sheet"], spec["cell"])
    return Result(
        value,
        [
            {
                "file": spec["file"],
                "sha256": src.file_sha(spec["file"]),
                "sheet": spec["sheet"],
                "range": spec["cell"],
            }
        ],
        f"cached value of cell {spec['cell']}",
    )


def _op_count(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    t = src.table(spec["table"])
    rows = select(t, spec.get("where"))
    distinct = spec.get("distinct")
    if distinct:
        _need(t, distinct)
        n = len({r.values.get(distinct) for r in rows if not is_blank(r.values.get(distinct))})
        how = f"number of distinct non-empty '{distinct}'"
        rng = t.locator_range(distinct, rows)
    else:
        n = len(rows)
        how = "number of data rows"
        rng = t.locator_range(None, rows)
    return Result(
        n, [_src(t, rng, rows)], f"{how} in table '{t.table_id}' where {spec.get('where')}"
    )


def _op_sum(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    t = src.table(spec["table"])
    col = spec["column"]
    _need(t, col)
    rows = select(t, spec.get("where"))
    nums = [to_number(r.values.get(col)) for r in rows]
    if any(n is None for n in nums):
        raise GoldError(f"sum: non-numeric '{col}' in {t.table_id}")
    return Result(
        round_num(sum(nums)),  # type: ignore[arg-type]
        [_src(t, t.locator_range(col, rows), rows)],
        f"sum of '{col}' over {len(rows)} row(s) of table '{t.table_id}' where {spec.get('where')}",
    )


def _op_row_sum(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Sum of several columns of the single row matching a filter."""
    t = src.table(spec["table"])
    cols = list(spec["columns"])
    _need(t, *cols)
    rows = select(t, spec.get("where"))
    if len(rows) != 1:
        raise GoldError(f"row_sum: expected exactly one row in {t.table_id}, matched {len(rows)}")
    nums = [to_number(rows[0].values.get(c)) for c in cols]
    if any(n is None for n in nums):
        raise GoldError(f"row_sum: non-numeric cell among {cols} in {t.table_id}")
    return Result(
        round_num(sum(nums)),  # type: ignore[arg-type]
        [_src(t, ",".join(rows[0].cells[c] for c in cols), rows)],
        f"sum of {cols} in the single row of table '{t.table_id}' matching {spec.get('where')}",
    )


def _op_values(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    t = src.table(spec["table"])
    col = spec["column"]
    _need(t, col)
    rows = select(t, spec.get("where"))
    vals = sorted({str(r.values[col]) for r in rows if not is_blank(r.values.get(col))})
    return Result(
        vals,
        [_src(t, t.locator_range(col, rows), rows)],
        f"distinct values of '{col}' in table '{t.table_id}' where {spec.get('where')}",
    )


def _op_group_count(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    t = src.table(spec["table"])
    by = list(spec["by"])
    distinct = spec.get("distinct")
    _need(t, *by, *([distinct] if distinct else []))
    rows = select(t, spec.get("where"))
    groups: dict[tuple, set | list] = {}
    skipped = 0
    for r in rows:
        k = tuple(r.values.get(b) for b in by)
        # a blank group key or a blank counted id is not a group / a unit
        if any(is_blank(x) for x in k) or (distinct and is_blank(r.values.get(distinct))):
            skipped += 1
            continue
        groups.setdefault(k, set() if distinct else [])
        if distinct:
            groups[k].add(r.values.get(distinct))
        else:
            groups[k].append(r)
    table_out = [
        {**{b: k[i] for i, b in enumerate(by)}, "n": len(v)}
        for k, v in sorted(groups.items(), key=lambda kv: tuple(str(x) for x in kv[0]))
    ]
    return Result(
        table_out,
        [_src(t, t.locator_range(b, rows), rows) for b in by],
        f"count of {'distinct ' + distinct if distinct else 'rows'} per {by} "
        f"in table '{t.table_id}' where {spec.get('where')}"
        + (f"; {skipped} row(s) with a blank group key or id excluded" if skipped else ""),
    )


def _op_group_stats(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Per-unit reduction, then group mean / sd / sem over units.

    ``unit`` names the unit of analysis (usually the animal). When several rows
    belong to one unit they are reduced with ``reduce`` first, so each unit
    contributes exactly one value to its group -- the point of the exercise.
    A unit whose rows disagree on a grouping column is an error, not a choice.
    """
    t = src.table(spec["table"])
    value = spec["value"]
    by = list(spec.get("by", []))
    unit_spec = spec.get("unit")
    units = [unit_spec] if isinstance(unit_spec, str) else list(unit_spec or [])
    rows = select(t, spec.get("where"))
    sources: list[dict[str, Any]] = []
    unmatched: list[str] = []
    if spec.get("join"):
        rows, unmatched = _attach(src, config, t, rows, spec["join"], sources)
        if unmatched and not spec["join"].get("allow_unmatched", False):
            raise GoldError(f"join left {len(unmatched)} unit(s) unmatched: {unmatched}")
    _need(t, value, *units)
    for c in by:
        if rows and c not in rows[0].values:
            raise KeyError(f"{t.table_id}: no column {c!r}")
    skipped: list[dict[str, Any]] = []
    per_unit: dict[Any, dict[str, Any]] = {}
    for r in rows:
        x = to_number(r.values.get(value))
        u = " / ".join(str(r.values.get(c)) for c in units) if units else f"row{r.locator}"
        if x is None:
            skipped.append({"unit": u, "row": r.locator, "value": r.values.get(value)})
            continue
        g = tuple(r.values.get(c) for c in by)
        slot = per_unit.setdefault(u, {"group": g, "xs": []})
        if slot["group"] != g:
            raise GoldError(f"unit {u!r} appears in two groups: {slot['group']} and {g}")
        slot["xs"].append(x)
    unit = " + ".join(units) if units else None
    reduce = spec.get("reduce", "single")
    groups: dict[tuple, list[float]] = {}
    for u, slot in per_unit.items():
        xs = slot["xs"]
        if reduce == "single":
            if len(xs) != 1:
                raise GoldError(f"unit {u!r} has {len(xs)} values; declare a reduce")
            v = xs[0]
        elif reduce == "sum":
            v = sum(xs)
        elif reduce == "mean":
            v = sum(xs) / len(xs)
        else:
            raise GoldError(f"unknown reduce {reduce!r}")
        groups.setdefault(slot["group"], []).append(v)
    out = [
        {**{b: k[i] for i, b in enumerate(by)}, **_stats(v)}
        for k, v in sorted(groups.items(), key=lambda kv: tuple(str(x) for x in kv[0]))
    ]
    used = [r for r in rows if to_number(r.values.get(value)) is not None]
    sources.insert(0, _src(t, t.locator_range(value, used), used))
    comp = (
        f"table '{t.table_id}' where {spec.get('where')}: "
        + (f"reduce '{value}' per unit '{unit}' by {reduce}, then " if unit else "")
        + f"mean, sd (n-1) and sem = sd/sqrt(n) of '{value}' per {by or 'whole table'}; "
        f"n counts units"
    )
    if spec.get("join"):
        j = spec["join"]
        comp += (
            f"; group columns {j['columns']} taken from table '{j['table']}' joined on "
            f"'{j['left_key']}'"
            + (f" (transform {j['left_transform']})" if j.get("left_transform") else "")
            + f" = '{j['right_key']}'"
            + (f" (transform {j['right_transform']})" if j.get("right_transform") else "")
        )
    if skipped:
        comp += f"; {len(skipped)} non-numeric cell(s) excluded: {skipped}"
    if unmatched:
        # allowed by the declaration (allow_unmatched), but never silent
        comp += (
            f"; {len(unmatched)} row(s) matched no right-table row and are outside the "
            f"selection ({len(set(unmatched))} distinct key(s))"
        )
    return Result(out, sources, comp)


def _keyed(
    src: RowSource, config: dict[str, Any], side: dict[str, Any]
) -> tuple[Table, dict[Any, Row]]:
    """Rows of one table keyed by unit, optionally reduced and value-mapped.

    ``reduce: sum`` collapses several rows of one unit into a synthetic row whose
    ``column`` holds the sum (the unit-of-analysis step). ``value_transform`` /
    ``value_map`` rewrite the compared value (e.g. a coded value 'A (label)' ->
    'A' -> 'label'); both are declared, never guessed.
    """
    t = src.table(side["table"])
    col = side["column"]
    _need(t, side["key"], col)
    tf = make_transform(config.get("transforms", {}).get(side.get("transform", "")))
    vt = make_transform(config.get("transforms", {}).get(side.get("value_transform", "")))
    vmap = side.get("value_map")
    grouped: dict[Any, list[Row]] = {}
    for r in select(t, side.get("where")):
        k = tf(r.values.get(side["key"]))
        if k is None:
            continue
        grouped.setdefault(k, []).append(r)
    out: dict[Any, Row] = {}
    for k, rs in grouped.items():
        if len(rs) > 1 and not side.get("reduce"):
            raise GoldError(f"key {k!r} repeats in {t.table_id}; declare a reduce")
        if side.get("reduce"):
            nums = [to_number(r.values.get(col)) for r in rs]
            if any(n is None for n in nums):
                raise GoldError(f"non-numeric '{col}' for unit {k!r} in {t.table_id}")
            if side["reduce"] == "sum":
                v: Any = round_num(sum(nums))  # type: ignore[arg-type]
            elif side["reduce"] == "mean":
                v = round_num(sum(nums) / len(nums))  # type: ignore[arg-type]
            else:
                raise GoldError(f"unknown reduce {side['reduce']!r}")
            cells = {col: t.locator_range(col, rs)}
            row = Row({**rs[0].values, col: v}, rs[0].locator, {**rs[0].cells, **cells})
        else:
            row = rs[0]
        if side.get("value_transform") or vmap:
            v2 = vt(row.values.get(col)) if side.get("value_transform") else row.values.get(col)
            if vmap is not None:
                v2 = vmap.get(str(v2), v2)
            row = Row({**row.values, col: v2}, row.locator, row.cells)
        out[k] = row
    return t, out


def _key_in(config: dict[str, Any], side: dict[str, Any], row: Row, keys: set[Any]) -> bool:
    tf = make_transform(config.get("transforms", {}).get(side.get("transform", "")))
    return tf(row.values.get(side["key"])) in keys


def _eq(a: Any, b: Any, tol: float) -> bool:
    x, y = to_number(a), to_number(b)
    if x is not None and y is not None:
        return abs(x - y) <= tol
    return norm_value(a) == norm_value(b)


def _op_compare(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Compare one attribute between two tables unit by unit."""
    ta, a = _keyed(src, config, spec["a"])
    tb, b = _keyed(src, config, spec["b"])
    tol = float(spec.get("tolerance", 0))
    common = sorted(set(a) & set(b), key=str)
    mismatches = []
    for k in common:
        va = a[k].values[spec["a"]["column"]]
        vb = b[k].values[spec["b"]["column"]]
        if not _eq(va, vb, tol):
            mismatches.append({"unit": k, "a": va, "b": vb})
    answer = {
        "units_compared": len(common),
        "all_equal": not mismatches,
        "mismatch_units": [str(m["unit"]) for m in mismatches],
        "mismatches": mismatches,
        "only_in_a": sorted(str(k) for k in set(a) - set(b)),
        "only_in_b": sorted(str(k) for k in set(b) - set(a)),
    }
    keep_a = set(common)
    rows_a = [
        r for r in select(ta, spec["a"].get("where")) if _key_in(config, spec["a"], r, keep_a)
    ]
    rows_b = [
        r for r in select(tb, spec["b"].get("where")) if _key_in(config, spec["b"], r, keep_a)
    ]
    return Result(
        answer,
        [
            _src(ta, ta.locator_range(spec["a"]["column"], rows_a), rows_a),
            _src(tb, tb.locator_range(spec["b"]["column"], rows_b), rows_b),
        ],
        f"'{spec['a']['column']}' of '{ta.table_id}' vs '{spec['b']['column']}' of "
        f"'{tb.table_id}' matched on unit key (tolerance {tol})",
    )


def _op_consistency(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Is one attribute (e.g. genotype) the same for a unit in every listed table?"""
    seen: dict[Any, list[tuple[str, Any, str]]] = {}
    sources = []
    for side in spec["sides"]:
        t = src.table(side["table"])
        _need(t, side["key"], side["column"])
        tf = make_transform(config.get("transforms", {}).get(side.get("transform", "")))
        used = []
        for r in select(t, side.get("where")):
            k = tf(r.values.get(side["key"]))
            if k is None:
                continue
            used.append(r)
            seen.setdefault(k, []).append(
                (t.table_id, r.values[side["column"]], r.cells[side["column"]])
            )
        sources.append(_src(t, t.locator_range(side["column"], used), used))
    conflicts = []
    for k in sorted(seen, key=str):
        distinct = {norm_value(v) for _, v, _ in seen[k]}
        if len(distinct) > 1:
            conflicts.append(
                {
                    "unit": k,
                    "values": [{"table": tid, "value": v, "cell": c} for tid, v, c in seen[k]],
                }
            )
    tables_per_unit = {k: {tid for tid, _, _ in v} for k, v in seen.items()}
    answer = {
        "units_checked": len(seen),
        "consistent": not conflicts,
        "conflict_units": [c["unit"] for c in conflicts],
        "conflicts": conflicts,
        "units_in_one_table_only": sorted(
            str(k) for k, v in tables_per_unit.items() if len(v) == 1
        ),
    }
    return Result(
        answer,
        sources,
        "attribute compared across "
        + ", ".join(f"{s['table']}.{s['column']}" for s in spec["sides"]),
    )


def _op_sum_check(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Does a recorded score equal the sum of the cells it claims to summarise?"""
    t = src.table(spec["table"])
    _need(t, spec["unit"], spec["recorded"], *spec["sum_of"])
    rows = select(t, spec.get("where"))
    mismatches = []
    checked = 0
    for r in rows:
        rec = to_number(r.values.get(spec["recorded"]))
        parts = [to_number(r.values.get(c)) for c in spec["sum_of"]]
        if rec is None or any(p is None for p in parts):
            continue
        checked += 1
        s = sum(parts)  # type: ignore[arg-type]
        if abs(rec - s) > 1e-9:
            mismatches.append(
                {"unit": r.values[spec["unit"]], "recorded": round_num(rec), "sum": round_num(s)}
            )
    return Result(
        {
            "rows_checked": checked,
            "mismatch_units": sorted(m["unit"] for m in mismatches),
            "mismatches": mismatches,
        },
        [
            _src(t, t.locator_range(spec["recorded"], rows), rows),
            *[_src(t, t.locator_range(c, rows), rows) for c in spec["sum_of"]],
        ],
        f"'{spec['recorded']}' compared with the sum of {spec['sum_of']} in table "
        f"'{t.table_id}'; rows with a non-numeric part are skipped",
    )


def _op_set_difference(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    a = evaluate(src, config, spec["a"])
    b = evaluate(src, config, spec["b"])
    diff = sorted(set(map(str, a.answer)) - set(map(str, b.answer)))
    return Result(diff, a.sources + b.sources, f"({a.computation}) minus ({b.computation})")


def _op_raw_range(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Non-empty values in one column of a raw sheet range, optionally regex-filtered."""
    file, sheet = spec["file"], spec["sheet"]
    col = spec["column"]
    lo, hi = int(spec["from_row"]), int(spec["to_row"])
    pat = re.compile(spec["regex"]) if spec.get("regex") else None
    vals = []
    for r in range(lo, hi + 1):
        v = src.raw_cell(file, sheet, f"{col}{r}")
        if v is None or v == "":
            continue
        if pat and not pat.fullmatch(str(v)):
            continue
        vals.append(str(v))
    return Result(
        sorted(set(vals)),
        [
            {
                "file": file,
                "sha256": src.file_sha(file),
                "sheet": sheet,
                "range": f"{col}{lo}:{col}{hi}",
            }
        ],
        f"distinct non-empty values of {col}{lo}:{col}{hi}"
        + (f" matching /{spec['regex']}/" if pat else ""),
    )


def _op_abstain(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Expected answer is an abstention; the evidence for it is checked, not asserted."""
    sources = []
    for check in spec.get("evidence", []):
        r = evaluate(src, config, check["op"])
        expect = check.get("expect")
        if "expect" in check and norm_value(r.answer) != expect:
            raise GoldError(
                f"abstention evidence no longer holds: {check['op']} gave {r.answer!r}, "
                f"expected {expect!r}"
            )
        sources.extend(r.sources)
    return Result(ABSTAIN, sources, spec.get("reason", "the data do not determine an answer"))


def _op_literal_checked(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Answer derived from a sub-op and mapped through a declared lookup."""
    r = evaluate(src, config, spec["of"])
    mapping = spec["map"]
    key = str(norm_value(r.answer))
    if key not in mapping:
        raise GoldError(f"no mapping for {key!r}")
    return Result(mapping[key], r.sources, r.computation + f"; mapped via {mapping}")


def _op_bundle(src: RowSource, config: dict[str, Any], spec: dict[str, Any]) -> Result:
    """Several named sub-answers, e.g. per-file counts for an inventory question."""
    out: dict[str, Any] = {}
    sources: list[dict[str, Any]] = []
    comps = []
    for name, sub in spec["parts"].items():
        r = evaluate(src, config, sub)
        out[name] = r.answer
        sources.extend(r.sources)
        comps.append(f"{name}: {r.computation}")
    return Result(out, sources, " || ".join(comps))


_OPS: dict[str, Callable[[RowSource, dict[str, Any], dict[str, Any]], Result]] = {
    "cell": _op_cell,
    "raw_cell": _op_raw_cell,
    "count": _op_count,
    "sum": _op_sum,
    "row_sum": _op_row_sum,
    "values": _op_values,
    "group_count": _op_group_count,
    "group_stats": _op_group_stats,
    "compare": _op_compare,
    "consistency": _op_consistency,
    "sum_check": _op_sum_check,
    "set_difference": _op_set_difference,
    "raw_range": _op_raw_range,
    "abstain": _op_abstain,
    "mapped": _op_literal_checked,
    "bundle": _op_bundle,
}

OPS = frozenset(_OPS)
