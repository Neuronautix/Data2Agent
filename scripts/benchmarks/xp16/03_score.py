"""Score an agent's answers against the XP16 gold questions.

Public and data-free: it reads a gold file and an answers file, and names
nothing about the dataset.

    python scripts/benchmarks/xp16/03_score.py --questions <pkg>/gold/questions.yaml \
        --answers <run>/answers.json [--package <pkg>] [--json score.json] \
        [--capabilities read_rows,filter,...]

Answers file (JSON)::

    {"answers": {
        "X16-Q06": {"answer": 18,
                    "sources": [{"file": "<path in source/>", "sha256": "<hex>",
                                 "sheet": "Sheet1", "cell": "E11"}]},
        "X16-Q35": {"abstain": true, "reason": "no sacrifice date recorded"}
    }}

A bare mapping ``{"X16-Q06": {...}}`` is accepted too. For delimited files a
source cites ``"line": <1-based physical line>`` and ``"field": <1-based>``.

Three things are scored and deliberately not folded into one number:

**Correctness.** exact / set / numeric-with-tolerance / table / object
comparison, as declared per question (``score:``) or derived from its
``answer_type``.

**Abstention.** A question whose gold is ``ABSTAIN`` or ``PENDING-<ids>`` has no
answer the data can support. Abstaining there is correct (+1). Answering it
confidently is the failure this benchmark exists to catch and costs a point
(-1), reported on its own line. Abstaining on an answerable question scores 0
and is counted as over-abstention -- safe, but not useful.

**Citation.** Each cited source must name a file of the frozen snapshot; its
sha256, when given, must match; the sheet and cell / range / line must exist;
and for a scalar answer citing one cell, that cell must hold the answer. A
source is ``verified`` only with a matching sha256 and an existing locator, and a
citation set takes the status of its weakest source (invalid < sha_unverified <
sha_only < verified); per-source statuses are kept in the report.

Per-category and per-capability breakdowns follow from the questions'
``category`` and ``requires`` tags. ``--capabilities`` restricts scoring to the
questions whose every requirement is in the given set -- the subset a given
build can be expected to answer.

Exit status: 0 when every scored question is correct or correctly abstained
and no confident answer was given to an unanswerable question; 1 otherwise.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

ABSTAIN = "ABSTAIN"
PENDING_PREFIX = "PENDING-"

# Replies that mean "I will not give a value". Matched on the whole normalised
# string, never as a substring: "unknown" is an abstention, "unknown genotype
# for animal X" is an answer.
_ABSTAIN_TOKENS = frozenset(
    {
        "abstain",
        "unknown",
        "cannot be determined",
        "can't be determined",
        "not determinable",
        "undetermined",
        "insufficient data",
        "not recorded",
        "no answer",
        "pending",
        "n/a",
    }
)

POINTS = {
    "correct": 1,
    "correct_abstention": 1,
    "incorrect": 0,
    "over_abstention": 0,
    "missing": 0,
    "confident_on_unanswerable": -1,
}


# ---------------------------------------------------------------- normalising


def _norm_str(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v)).strip().casefold()


def _num(v: Any) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(float(v)) else None
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def is_unanswerable(expected: Any) -> bool:
    return expected == ABSTAIN or (
        isinstance(expected, str) and expected.startswith(PENDING_PREFIX)
    )


def is_abstention(entry: dict[str, Any] | None) -> bool:
    if not isinstance(entry, dict):
        return False
    if entry.get("abstain") is True:
        return True
    ans = entry.get("answer")
    if ans is None and "answer" in entry:
        return True
    if isinstance(ans, str):
        s = _norm_str(ans)
        return s in _ABSTAIN_TOKENS or s.startswith(_norm_str(PENDING_PREFIX))
    return False


# ---------------------------------------------------------------- comparison


def default_rule(question: dict[str, Any]) -> dict[str, Any]:
    if question.get("score"):
        return question["score"]
    t = question.get("answer_type")
    if t == "number":
        tol = question.get("tolerance", {"abs": 0.001})
        return {"type": "number", **tol}
    if t in {"integer", "boolean", "set", "string", "exact"}:
        return {"type": t}
    return {"type": "exact"}


def compare(expected: Any, got: Any, rule: dict[str, Any]) -> tuple[bool, str]:
    """Compare one value under one rule. Returns (ok, human-readable reason)."""
    kind = rule.get("type", "exact")
    if kind == "number":
        e, g = _num(expected), _num(got)
        if e is None or g is None:
            return False, f"not numeric: expected {expected!r}, got {got!r}"
        tol = float(rule.get("abs", 0.0))
        rel = rule.get("rel")
        limit = max(tol, abs(e) * float(rel)) if rel is not None else tol
        ok = abs(e - g) <= limit + 1e-12
        return ok, "" if ok else f"{g} differs from {e} by more than {limit}"
    if kind == "integer":
        e, g = _num(expected), _num(got)
        ok = e is not None and g is not None and e == g
        return ok, "" if ok else f"expected {expected!r}, got {got!r}"
    if kind == "boolean":
        g = got
        if isinstance(got, str) and _norm_str(got) in {"true", "yes"}:
            g = True
        elif isinstance(got, str) and _norm_str(got) in {"false", "no"}:
            g = False
        ok = isinstance(g, bool) and g == expected
        return ok, "" if ok else f"expected {expected!r}, got {got!r}"
    if kind in {"exact", "string"}:
        e, g = _num(expected), _num(got)
        if e is not None and g is not None:
            ok = e == g
        else:
            ok = _norm_str(expected) == _norm_str(got)
        return ok, "" if ok else f"expected {expected!r}, got {got!r}"
    if kind == "set":
        if not isinstance(got, (list, tuple, set)):
            return False, f"expected a list, got {type(got).__name__}"
        es = {_norm_str(x) for x in expected}
        gs = {_norm_str(x) for x in got}
        if es == gs:
            return True, ""
        return False, f"missing {sorted(es - gs)}, unexpected {sorted(gs - es)}"
    if kind == "table":
        return _compare_table(expected, got, rule)
    if kind == "object":
        if not isinstance(got, dict):
            return False, f"expected an object, got {type(got).__name__}"
        problems = []
        for name, sub in rule.get("fields", {}).items():
            if name not in got:
                problems.append(f"{name}: missing")
                continue
            ok, why = compare(expected.get(name), got[name], sub)
            if not ok:
                problems.append(f"{name}: {why}")
        return not problems, "; ".join(problems)
    raise ValueError(f"unknown comparison type {kind!r}")


def _compare_table(expected: Any, got: Any, rule: dict[str, Any]) -> tuple[bool, str]:
    if not isinstance(got, list) or not all(isinstance(r, dict) for r in got):
        return False, "expected a list of row objects"
    keys = rule.get("key", [])

    def k(row: dict[str, Any]) -> tuple[str, ...]:
        return tuple(_norm_str(row.get(c)) for c in keys)

    exp_rows = {k(r): r for r in expected}
    got_rows: dict[tuple[str, ...], dict[str, Any]] = {}
    for r in got:
        if k(r) in got_rows:
            return False, f"duplicate row for key {k(r)}"
        got_rows[k(r)] = r
    problems = []
    for key, er in exp_rows.items():
        gr = got_rows.get(key)
        if gr is None:
            problems.append(f"row {key}: missing")
            continue
        for name, sub in rule.get("fields", {}).items():
            if name not in gr:
                problems.append(f"row {key}: {name} missing")
                continue
            ok, why = compare(er.get(name), gr[name], sub)
            if not ok:
                problems.append(f"row {key}: {name}: {why}")
    extra = sorted(set(got_rows) - set(exp_rows))
    if extra:
        problems.append(f"unexpected row(s) {extra}")
    return not problems, "; ".join(problems)


# ---------------------------------------------------------------- citations


class CitationContext:
    """What a citation can be checked against: checksums, and optionally the bytes."""

    def __init__(self, checksums: dict[str, Any] | None, source_dir: Path | None) -> None:
        self.shas = {f["path"]: f["sha256"] for f in (checksums or {}).get("files", [])}
        self.source_dir = source_dir
        self._books: dict[str, Any] = {}
        self._lines: dict[str, list[list[str]]] = {}
        self._boris: dict[str, Any] = {}

    def _book(self, file: str) -> Any | None:
        if self.source_dir is None:
            return None
        if file not in self._books:
            try:
                import openpyxl  # noqa: PLC0415
            except ImportError:
                self._books[file] = None
                return None
            data = (self.source_dir / file).read_bytes()
            try:
                self._books[file] = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            except Exception:  # noqa: BLE001 -- not a workbook: treated as delimited
                self._books[file] = None
        return self._books[file]

    def _delimited(self, file: str) -> list[list[str]] | None:
        if self.source_dir is None:
            return None
        if file not in self._lines:
            import csv  # noqa: PLC0415

            text = (self.source_dir / file).read_bytes().decode("utf-8", errors="replace")
            delim = "\t" if file.lower().endswith(".tsv") else ","
            self._lines[file] = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
        return self._lines[file]

    def check(self, src: dict[str, Any], answer: Any, scalar: bool) -> dict[str, Any]:
        file = src.get("file")
        out: dict[str, Any] = {"file": file}
        if file not in self.shas:
            out["status"] = "invalid"
            out["reason"] = "file is not part of the frozen snapshot"
            return out
        sha = src.get("sha256")
        out["sha256_match"] = None if sha is None else (sha == self.shas[file])
        if sha is not None and sha != self.shas[file]:
            out["status"] = "invalid"
            out["reason"] = "sha256 does not match the frozen file"
            return out
        locator_ok, value = self._locate(file, src)
        out["locator_exists"] = locator_ok
        if locator_ok is False:
            out["status"] = "invalid"
            out["reason"] = "cited sheet / cell / line does not exist"
            return out
        if scalar and value is not _NO_VALUE:
            numeric = _num(value) is not None and _num(answer) is not None
            rule = {"type": "number", "abs": 1e-6} if numeric else {"type": "exact"}
            ok, _ = compare(value, answer, rule)
            out["supports_answer"] = ok
            if not ok:
                out["status"] = "invalid"
                out["reason"] = f"cited cell holds {value!r}, not the answer"
                return out
        if sha is None:
            out["status"] = "sha_unverified"
        elif locator_ok is True:
            out["status"] = "verified"
        else:  # sha matches, but the locator could not be checked
            out["status"] = "sha_only"
        return out

    def _boris_locator(self, file: str, src: dict[str, Any]) -> bool | None:
        """A BORIS row as the service locates it: observation id + event index(es).

        The observation must exist in the project and every cited event index
        must be one of its events (0-based, as the service reports them).
        """
        if self.source_dir is None:
            return None
        if file not in self._boris:
            try:
                self._boris[file] = json.loads(
                    (self.source_dir / file).read_text(encoding="utf-8")
                ).get("observations", {})
            except (OSError, ValueError):
                self._boris[file] = None
        observations = self._boris[file]
        if observations is None or src["observation_id"] not in observations:
            return False
        n = len(observations[src["observation_id"]].get("events", []))
        indices = [
            src[k]
            for k in ("event_index", "start_event_index", "stop_event_index")
            if src.get(k) is not None
        ]
        return all(isinstance(i, int) and 0 <= i < n for i in indices)

    def _locate(self, file: str, src: dict[str, Any]) -> tuple[bool | None, Any]:
        """(exists?, value of a single cited cell or _NO_VALUE). None = not checkable."""
        if src.get("observation_id") is not None:
            return self._boris_locator(file, src), _NO_VALUE
        book = self._book(file) if src.get("sheet") is not None else None
        if src.get("sheet") is not None:
            if book is None:
                return (None, _NO_VALUE) if self.source_dir is None else (False, _NO_VALUE)
            if src["sheet"] not in book.sheetnames:
                return False, _NO_VALUE
            ws = book[src["sheet"]]
            loc = src.get("cell") or src.get("range")
            if loc is None and src.get("row") is not None:
                return int(src["row"]) <= ws.max_row, _NO_VALUE
            if loc is None:
                return True, _NO_VALUE
            cells = re.findall(r"([A-Z]{1,3})(\d+)", str(loc))
            if not cells:
                return False, _NO_VALUE
            from openpyxl.utils import column_index_from_string  # noqa: PLC0415

            for col, row in cells:
                if int(row) > ws.max_row or column_index_from_string(col) > ws.max_column:
                    return False, _NO_VALUE
            if len(cells) == 1 and ":" not in str(loc):
                return True, ws[f"{cells[0][0]}{cells[0][1]}"].value
            return True, _NO_VALUE
        # 'line 6, field 4', 'lines 4-12, field 4', or a composed column
        # 'lines 4-12, field 1+3' (a key built from fields 1 and 3)
        m = re.fullmatch(
            r"lines? (\d+)(?:-(\d+))?(?:, field (\d+(?:\+\d+)*))?",
            str(src.get("range", "")).strip(),
        )
        if m and src.get("line") is None:
            lo, hi, fld = m.group(1), m.group(2), m.group(3)
            fields = [int(f) for f in fld.split("+")] if fld else []
            if hi is None and len(fields) <= 1:
                src = {**src, "line": int(lo), **({"field": fields[0]} if fields else {})}
            else:
                lines = self._delimited(file)
                if lines is None:
                    return None, _NO_VALUE
                lo_i, hi_i = int(lo), int(hi or lo)
                if not 1 <= lo_i <= hi_i <= len(lines):
                    return False, _NO_VALUE
                width = max((len(r) for r in lines[lo_i - 1 : hi_i]), default=0)
                return all(1 <= f <= width for f in fields), _NO_VALUE
        if src.get("line") is not None:
            lines = self._delimited(file)
            if lines is None:
                return None, _NO_VALUE
            n = int(src["line"])
            if n < 1 or n > len(lines):
                return False, _NO_VALUE
            if src.get("field") is not None:
                f = int(src["field"])
                if f < 1 or f > len(lines[n - 1]):
                    return False, _NO_VALUE
                return True, lines[n - 1][f - 1]
            return True, _NO_VALUE
        return None, _NO_VALUE


_NO_VALUE = object()

# Per-source citation statuses, weakest first. The status of a citation set is
# the weakest of its sources.
#   invalid         file not in the snapshot, sha256 mismatch, missing locator,
#                   or a cited cell that does not hold the answer
#   sha_unverified  no sha256 given, so the bytes behind it are not pinned
#   sha_only        sha256 matches, locator could not be checked
#   verified        sha256 matches and the locator exists (and supports a scalar)
CITATION_STRENGTH = {"invalid": 0, "sha_unverified": 1, "sha_only": 2, "verified": 3}


def citation_status(
    entry: dict[str, Any], question: dict[str, Any], ctx: CitationContext | None
) -> dict[str, Any]:
    sources = entry.get("sources") if isinstance(entry, dict) else None
    if not sources:
        return {"status": "missing", "checks": []}
    if ctx is None:
        return {"status": "unchecked", "checks": []}
    scalar = question.get("answer_type") in {"number", "integer", "string", "exact"}
    checks = [ctx.check(s, entry.get("answer"), scalar and len(sources) == 1) for s in sources]
    # A citation set is only as strong as its weakest source: one verified source
    # must not vouch for another that could not be checked.
    status = min((c["status"] for c in checks), key=CITATION_STRENGTH.__getitem__)
    gold_files = {s.get("file") for s in question.get("sources", [])}
    return {
        "status": status,
        "per_source": [c["status"] for c in checks],
        "checks": checks,
        "cites_a_gold_source_file": any(s.get("file") in gold_files for s in sources),
    }


# ---------------------------------------------------------------- scoring


def score_question(
    q: dict[str, Any], entry: dict[str, Any] | None, ctx: CitationContext | None
) -> dict[str, Any]:
    expected = q.get("expected")
    unanswerable = is_unanswerable(expected)
    out: dict[str, Any] = {
        "id": q["id"],
        "category": q.get("category"),
        "requires": q.get("requires", []),
    }
    if entry is None:
        out["outcome"] = "missing"
    elif is_abstention(entry):
        out["outcome"] = "correct_abstention" if unanswerable else "over_abstention"
    elif unanswerable:
        out["outcome"] = "confident_on_unanswerable"
        out["detail"] = f"gold is {expected}; an answer was given"
    else:
        ok, why = compare(expected, entry.get("answer"), default_rule(q))
        out["outcome"] = "correct" if ok else "incorrect"
        if why:
            out["detail"] = why
    out["points"] = POINTS[out["outcome"]]
    if entry is not None and not is_abstention(entry):
        out["citation"] = citation_status(entry, q, ctx)
    return out


def load_answers(path: Path) -> dict[str, Any]:
    doc = json.loads(path.read_text(encoding="utf-8"))
    answers = doc.get("answers", doc) if isinstance(doc, dict) else {}
    if not isinstance(answers, dict):
        raise ValueError("answers must be an object keyed by question id")
    return answers


def select_questions(
    questions: list[dict[str, Any]], capabilities: set[str] | None
) -> list[dict[str, Any]]:
    if capabilities is None:
        return list(questions)
    return [q for q in questions if set(q.get("requires", [])) <= capabilities]


def score_run(
    questions: list[dict[str, Any]],
    answers: dict[str, Any],
    ctx: CitationContext | None = None,
    capabilities: set[str] | None = None,
) -> dict[str, Any]:
    chosen = select_questions(questions, capabilities)
    results = [score_question(q, answers.get(q["id"]), ctx) for q in chosen]
    unknown_ids = sorted(set(answers) - {q["id"] for q in questions})

    def tally(rows: list[dict[str, Any]]) -> dict[str, Any]:
        counts: dict[str, int] = {}
        for r in rows:
            counts[r["outcome"]] = counts.get(r["outcome"], 0) + 1
        return {
            "n": len(rows),
            "points": sum(r["points"] for r in rows),
            **dict(sorted(counts.items())),
        }

    by_category: dict[str, list[dict[str, Any]]] = {}
    by_capability: dict[str, list[dict[str, Any]]] = {}
    for r in results:
        by_category.setdefault(r["category"] or "?", []).append(r)
        for cap in r["requires"]:
            by_capability.setdefault(cap, []).append(r)
    citations: dict[str, int] = {}
    for r in results:
        if "citation" in r:
            s = r["citation"]["status"]
            citations[s] = citations.get(s, 0) + 1
    return {
        "scored": len(results),
        "excluded_by_capability_filter": len(questions) - len(chosen),
        "total": tally(results),
        "by_category": {k: tally(v) for k, v in sorted(by_category.items())},
        "by_capability": {k: tally(v) for k, v in sorted(by_capability.items())},
        "citations": dict(sorted(citations.items())),
        "confident_on_unanswerable": [
            r["id"] for r in results if r["outcome"] == "confident_on_unanswerable"
        ],
        "answers_for_unknown_ids": unknown_ids,
        "results": results,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--questions", type=Path, required=True)
    parser.add_argument("--answers", type=Path, required=True)
    parser.add_argument("--package", type=Path, help="benchmark package, for citation checks")
    parser.add_argument("--json", type=Path, help="write the full report here")
    parser.add_argument("--capabilities", help="comma-separated; score only covered questions")
    args = parser.parse_args()

    import yaml  # noqa: PLC0415

    gold = yaml.safe_load(args.questions.read_text(encoding="utf-8"))
    ctx = None
    if args.package:
        checksums = json.loads((args.package / "checksums.json").read_text(encoding="utf-8"))
        if checksums.get("dataset_id") != gold.get("dataset_id"):
            sys.exit("gold and package disagree on dataset_id; refusing to score")
        ctx = CitationContext(checksums, args.package / "source")
    caps = set(args.capabilities.split(",")) if args.capabilities else None
    report = score_run(gold["questions"], load_answers(args.answers), ctx, caps)
    report["dataset_id"] = gold.get("dataset_id")

    t = report["total"]
    print(f"scored {report['scored']} question(s); points {t['points']} / {report['scored']}")
    for k in ("correct", "correct_abstention", "incorrect", "over_abstention", "missing"):
        print(f"  {k:<26} {t.get(k, 0)}")
    bad = report["confident_on_unanswerable"]
    print(f"  CONFIDENT ON UNANSWERABLE  {len(bad)}  {bad if bad else ''}")
    print("by capability:")
    for cap, c in report["by_capability"].items():
        print(f"  {cap:<24} {c.get('correct', 0) + c.get('correct_abstention', 0)}/{c['n']}")
    if report["citations"]:
        print("citations:", report["citations"])
    if args.json:
        args.json.write_text(json.dumps(report, indent=2, default=str) + "\n", encoding="utf-8")
    good = t.get("correct", 0) + t.get("correct_abstention", 0)
    return 0 if good == report["scored"] and not bad else 1


if __name__ == "__main__":
    raise SystemExit(main())
