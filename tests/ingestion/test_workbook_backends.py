"""Workbook formats beyond OOXML, behind one table contract (D2A-94).

The fixtures in ``tests/fixtures/workbooks`` are one workbook saved by Microsoft
Excel itself as .xlsx, .xls (BIFF8), .xlsb and .ods; ``make_fixtures.ps1`` is
the record of how. Testing against bytes Excel wrote -- rather than bytes a
Python library believes a format looks like -- is the point: a reader that
round-trips its own output proves nothing about the files a lab actually has.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest

from data2agent.ingest import formats
from data2agent.ingest.pipeline import ingest
from data2agent.readers import workbook

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "workbooks"
FORMATS = ("xlsx", "xls", "xlsb", "ods")


@pytest.mark.parametrize("extension", FORMATS)
def test_each_workbook_format_is_identified_from_its_bytes(extension: str):
    detected, notes = formats.detect(FIXTURES / f"animals.{extension}")
    assert detected.format_id == extension
    assert detected.detected_by == "container"
    assert detected.extension_conflict is False
    assert detected.format_id in formats.WORKBOOK_FORMATS
    assert notes == []


def test_a_real_xls_under_the_wrong_name_is_identified_by_content(tmp_path: Path):
    """The inverse of XP14's case: BIFF bytes wearing an OOXML name."""
    renamed = tmp_path / "results.xlsx"
    shutil.copyfile(FIXTURES / "animals.xls", renamed)

    detected, notes = formats.detect(renamed)
    assert detected.format_id == "xls"
    assert detected.extension_format == "xlsx"
    assert detected.extension_conflict is True
    assert notes and "'xls'" in notes[0]


def test_an_ole2_file_without_a_workbook_stream_is_not_called_xls(tmp_path: Path):
    """The OLE2 signature alone is shared by .doc, .ppt and .msg; it is not an .xls."""
    fake = tmp_path / "fake.xls"
    fake.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 1024)

    detected, notes = formats.detect(fake)
    assert detected.format_id == "ole2-container"
    assert detected.extension_conflict is True
    assert detected.format_id not in formats.WORKBOOK_FORMATS
    assert notes


_NOSTREAM = 0xFFFFFFFF


def _compound_file(entries: list[tuple[str, int, int, int, int]]) -> bytes:
    """A minimal [MS-CFB] v3 file: header, one FAT sector, one directory sector.

    ``entries`` are ``(name, type, left, right, child)``; type 5 is the root
    storage, 1 a storage, 2 a stream. Streams are empty -- only the directory
    is under test, and detection never reads a stream.
    """
    import struct

    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HHHHH", header, 0x18, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<I", header, 0x2C, 1)  # one FAT sector
    struct.pack_into("<I", header, 0x30, 1)  # directory starts at sector 1
    struct.pack_into("<III", header, 0x38, 4096, 0xFFFFFFFE, 0)
    struct.pack_into("<II", header, 0x44, 0xFFFFFFFE, 0)
    struct.pack_into("<109I", header, 0x4C, 0, *([_NOSTREAM] * 108))

    fat = struct.pack("<128I", 0xFFFFFFFD, 0xFFFFFFFE, *([_NOSTREAM] * 126))
    directory = bytearray(512)
    for index, (name, kind, left, right, child) in enumerate(entries):
        encoded = name.encode("utf-16-le") + b"\x00\x00"
        base = index * 128
        directory[base : base + len(encoded)] = encoded
        struct.pack_into("<HBB", directory, base + 0x40, len(encoded), kind, 1)
        struct.pack_into("<III", directory, base + 0x44, left, right, child)
    return bytes(header) + fat + bytes(directory)


def test_a_workbook_stream_at_the_root_makes_an_ole2_file_xls(tmp_path: Path):
    path = tmp_path / "minimal.xls"
    path.write_bytes(
        _compound_file(
            [
                ("Root Entry", 5, _NOSTREAM, _NOSTREAM, 1),
                ("Workbook", 2, _NOSTREAM, _NOSTREAM, _NOSTREAM),
            ]
        )
    )
    assert formats.detect(path)[0].format_id == "xls"


def test_an_embedded_workbook_does_not_make_the_outer_document_xls(tmp_path: Path):
    """A Word file embedding an Excel object has a nested 'Workbook' stream."""
    path = tmp_path / "report.doc"
    path.write_bytes(
        _compound_file(
            [
                ("Root Entry", 5, _NOSTREAM, _NOSTREAM, 1),
                ("WordDocument", 2, _NOSTREAM, 2, _NOSTREAM),
                ("ObjectPool", 1, _NOSTREAM, _NOSTREAM, 3),
                ("Workbook", 2, _NOSTREAM, _NOSTREAM, _NOSTREAM),
            ]
        )
    )
    detected, _ = formats.detect(path)
    assert detected.format_id == "ole2-container"
    assert detected.format_id not in formats.WORKBOOK_FORMATS


def test_a_zip_holding_another_opendocument_type_is_not_ods(tmp_path: Path):
    import zipfile

    text = tmp_path / "notes.ods"
    with zipfile.ZipFile(text, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", "<office:document-content/>")

    detected, _ = formats.detect(text)
    assert detected.format_id == "zip-container"
    assert detected.extension_conflict is True


# -- reading: needs both optional parsers --------------------------------------

readers = pytest.mark.skipif(
    not (workbook.OPENPYXL.available() and workbook.CALAMINE.available()),
    reason="needs the 'xlsx' and 'workbooks' extras",
)


@readers
def test_backends_agree_on_where_a_sheet_starts(tmp_path: Path):
    """A used range starting at C3 keeps column A as position 0 in both backends."""
    import openpyxl

    path = tmp_path / "offset.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["C3"], sheet["D3"], sheet["C4"], sheet["D4"], sheet["B6"] = "id", "x", 1, 2, "late"
    book.save(path)

    rows = {}
    for backend in (workbook.OPENPYXL, workbook.CALAMINE):
        with workbook.open_workbook(path, backend) as opened:
            rows[backend.name] = [tuple(row) for row in opened.iter_rows("Sheet")]
    assert rows["calamine"] == rows["openpyxl"]
    assert rows["calamine"][2] == (None, None, "id", "x")


def _contract(sheet) -> dict:
    """A sheet profile minus the fields that legitimately name the file or backend."""
    payload = sheet.as_dict()
    for key in ("path", "workbook", "reader"):
        payload.pop(key)
    return payload


@readers
def test_an_empty_sheet_is_profiled_as_empty_by_calamine(tmp_path: Path):
    """calamine panics iterating a sheet with no used range; XP1's .xlsb exports have one."""
    import openpyxl

    path = tmp_path / "with-empty.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["id"])
    book.active.append([1])
    book.create_sheet("empty")
    book.save(path)

    with workbook.open_workbook(path, workbook.CALAMINE) as opened:
        assert list(opened.iter_rows("empty")) == []
        sheet = workbook._profile_sheet(
            opened, "empty", "with-empty.xlsx", 1, workbook.DEFAULT_CONVENTION
        )
    assert sheet.profiled is True
    assert sheet.rows == 0
    assert any("empty" in warning for warning in sheet.warnings)


def test_a_parser_panic_becomes_an_unreadable_sheet_not_an_aborted_ingest():
    class PanicException(BaseException):  # the shape pyo3 raises
        pass

    with pytest.raises(RuntimeError, match="could not parse"):
        with workbook._panics_as_errors():
            raise PanicException("called `Option::unwrap()` on a `None` value")
    with pytest.raises(KeyboardInterrupt):
        with workbook._panics_as_errors():
            raise KeyboardInterrupt


@readers
def test_one_workbook_profiles_identically_in_every_format():
    """Same high-level table contract regardless of backend."""
    profiles = {
        extension: workbook.profile_workbook(
            FIXTURES / f"animals.{extension}", "animals", format_id=extension
        )
        for extension in FORMATS
    }
    reference = [_contract(sheet) for sheet in profiles["xlsx"]]
    for extension in FORMATS:
        assert [_contract(sheet) for sheet in profiles[extension]] == reference, extension

    animals, notes = profiles["xls"]
    assert animals.reader is workbook.CALAMINE
    assert profiles["xlsx"][0].reader is workbook.OPENPYXL
    assert animals.header_row == 2  # row 1 is blank in the source
    assert animals.rows == 3
    assert [c.name for c in animals.columns] == [
        "animal_id",
        "genotype",
        "weight_g",
        "dob",
        "score",
        "note",
    ]
    assert notes.sheet_state == "hidden"


@readers
def test_a_real_xls_dataset_is_ingested_and_its_rows_are_readable(tmp_path: Path):
    from data2agent.mcp.service import DatasetService

    source = tmp_path / "ds"
    source.mkdir()
    shutil.copyfile(FIXTURES / "animals.xls", source / "animals.xls")
    out = tmp_path / "out"
    result = ingest(source, out)

    sheet = result.manifest["tables"]["animals.xls#animals"]
    assert sheet["profiled"] is True
    assert sheet["reader"] == {"backend": "calamine", "cell_values": "cached"}
    assert result.manifest["tables"]["animals.xls#notes"]["sheet_state"] == "hidden"
    readers = result.provenance["configuration"]["workbook_readers"]
    assert readers["calamine"] == workbook.CALAMINE.version()

    service = DatasetService(out, mode="structured")
    payload = service.read_rows("animals.xls#animals", offset=1, limit=2)
    first, second = payload["rows"]
    # source_row is the spreadsheet's own row number, whatever the backend.
    assert first["source_row"] == 4
    assert first["values"]["animal_id"] == "A002"
    assert first["values"]["weight_g"] == 19
    assert first["values"]["dob"] == "2024-01-04T00:00:00"
    assert first["missing"]["note"] == {"kind": "empty"}
    assert second["values"]["genotype"] is None
    assert second["missing"]["genotype"] == {"kind": "sentinel", "raw": "NA"}


@readers
@pytest.mark.parametrize("extension", FORMATS)
def test_rows_agree_across_formats(tmp_path: Path, extension: str):
    """A formula is read as its cached value (C3*2 = 43) by both backends."""
    from data2agent.mcp.service import DatasetService

    source = tmp_path / "ds"
    source.mkdir()
    shutil.copyfile(FIXTURES / f"animals.{extension}", source / f"animals.{extension}")
    ingest(source, tmp_path / "out")
    service = DatasetService(tmp_path / "out", mode="structured")

    rows = service.read_rows(f"animals.{extension}#animals")["rows"]
    assert [row["source_row"] for row in rows] == [3, 4, 5]
    assert rows[0]["values"] == {
        "animal_id": "A001",
        "genotype": "WT",
        "weight_g": 21.5,
        "dob": "2024-01-03T00:00:00",
        "score": 43,
        "note": "ok",
    }


@readers
def test_an_absent_backend_leaves_the_workbook_unprofiled_not_empty(tmp_path: Path, monkeypatch):
    """Missing calamine must not turn a real .xls into a dataset with no tables."""
    monkeypatch.setattr(
        workbook.Backend, "available", lambda self: self.name != workbook.CALAMINE.name
    )
    source = tmp_path / "ds"
    source.mkdir()
    shutil.copyfile(FIXTURES / "animals.xls", source / "animals.xls")
    shutil.copyfile(FIXTURES / "animals.xlsx", source / "animals.xlsx")
    out = tmp_path / "out"
    result = ingest(source, out)

    tables = result.manifest["tables"]
    assert not any(key.startswith("animals.xls#") for key in tables)
    assert "animals.xlsx#animals" in tables  # openpyxl is unaffected
    assert any(
        "animals.xls" in warning and "'workbooks' extra" in warning
        for warning in result.manifest["warnings"]
    )
    candidates = {item["path"]: item for item in result.manifest["metadata_candidates"]}
    assert candidates["animals.xls"]["reason"] == "reader-unavailable"

    evidence = json.loads((out / "evidence.json").read_text(encoding="utf-8"))
    unavailable = [
        item
        for claim in evidence["claims"]
        for item in claim["evidence"]
        if item["check"] == "workbook.reader-unavailable"
    ]
    assert unavailable[0]["result"] == {
        "format": "xls",
        "extra": "workbooks",
        "backend": "calamine",
    }
    assert "calamine" not in result.provenance["configuration"]["workbook_readers"]


@readers
def test_every_format_validates_against_the_published_schema(tmp_path: Path):
    jsonschema = pytest.importorskip("jsonschema")
    source = tmp_path / "ds"
    source.mkdir()
    for extension in FORMATS:
        shutil.copyfile(FIXTURES / f"animals.{extension}", source / f"animals.{extension}")
    out = tmp_path / "out"
    ingest(source, out)

    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    schema = json.loads(Path("schemas/dataset-manifest.schema.json").read_text(encoding="utf-8"))
    jsonschema.validate(manifest, schema)
    assert len(manifest["tables"]) == 2 * len(FORMATS)


@readers
def test_the_source_workbook_is_never_modified(tmp_path: Path):
    from data2agent.ingest.checksum import hash_file

    source = tmp_path / "ds"
    source.mkdir()
    for extension in ("xls", "xlsb", "ods"):
        shutil.copyfile(FIXTURES / f"animals.{extension}", source / f"animals.{extension}")
    before = {path.name: hash_file(path) for path in source.iterdir()}
    ingest(source, tmp_path / "out")
    assert {path.name: hash_file(path) for path in source.iterdir()} == before
    assert sorted(path.name for path in source.iterdir()) == sorted(before)
