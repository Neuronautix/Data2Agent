"""What content-recognised metadata does to the five indicators that cascade from it.

F2, F3, F4, I2 and R1.2 all read ``manifest.metadata_files``. Before D2A-49a they
read a list that only a filename could get onto, so a dataset whose metadata sat
in a spreadsheet was assessed as a dataset with no metadata. These tests pin down
what changes, and -- as importantly -- what a check is still not allowed to say
about bytes it never read.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

pytest.importorskip("yaml", reason="profiles need the 'fair' extra")

from data2agent.ingest.pipeline import ingest  # noqa: E402
from data2agent.mcp import DatasetService  # noqa: E402


def _registry_rows(count: int = 12) -> str:
    header = "animal_id,batch,sex,group\n"
    body = "".join(
        f"A{index:03d},{index % 3 + 1},{'MF'[index % 2]},{'control' if index % 2 else 'treated'}\n"
        for index in range(count)
    )
    return header + body


def _measurement_rows(animals: int = 12) -> str:
    header = "animal_id,session,latency_s\n"
    return header + "".join(
        f"A{animal:03d},{session},{12.5 + animal * 0.25 + session}\n"
        for animal in range(animals)
        for session in (1, 2, 3)
    )


def _assess(tmp_path: Path, files: dict[str, str], name: str = "run") -> dict[str, dict]:
    root = tmp_path / name / "dataset"
    root.mkdir(parents=True)
    for filename, text in files.items():
        (root / filename).write_text(text, encoding="utf-8")
    result = ingest(root, tmp_path / name / "out")
    service = DatasetService(result.output_dir, mode="fair-deterministic")
    return {item["rule_id"]: item for item in service.run_fair_check()["results"]}


def test_a_registry_table_makes_the_dataset_described(tmp_path: Path):
    """The verdict that changes: F2 was a fail on exactly these bytes before."""
    with_registry = _assess(
        tmp_path,
        {"IDs Batch Sex Group.csv": _registry_rows(), "trials.csv": _measurement_rows()},
        name="with",
    )
    without_registry = _assess(tmp_path, {"trials.csv": _measurement_rows()}, name="without")

    assert with_registry["F2-METADATA-PRESENT"]["result"] == "pass"
    assert without_registry["F2-METADATA-PRESENT"]["result"] == "fail"


def test_the_recognition_basis_travels_into_the_verdict(tmp_path: Path):
    results = _assess(tmp_path, {"registry.csv": _registry_rows(), "t.csv": _measurement_rows()})
    assert results["F2-METADATA-PRESENT"]["result"] == "pass"
    assert results["F2-METADATA-PRESENT"]["evidence"], "a pass must cite what it read"


def test_a_failing_f2_no_longer_claims_recognition_is_by_filename_alone(tmp_path: Path):
    results = _assess(tmp_path, {"trials.csv": _measurement_rows()})
    rationale = results["F2-METADATA-PRESENT"]["rationale"]
    assert "no filename matched a known convention" in rationale
    assert "content" in rationale


def test_an_embedded_recognition_leaves_the_file_a_data_file(tmp_path: Path):
    """A registry CSV is metadata and payload. F3 must still expect it to be referenced."""
    results = _assess(
        tmp_path,
        {
            "README.md": "# Study\n\nSee trials.csv.\n",
            "registry.csv": _registry_rows(),
            "trials.csv": _measurement_rows(),
        },
    )
    linkage = results["F3-METADATA-LINKS-DATA"]
    assert linkage["result"] == "fail"
    assert "registry.csv" in linkage["rationale"], (
        "an embedded recognition must not remove its file from the data files"
    )


def test_a_json_descriptor_under_an_unconventional_name_declares_provenance(tmp_path: Path):
    """R1.2 changes verdict: the descriptor was invisible when only names were read."""
    document = {
        "@context": "https://schema.org/",
        "@type": "Dataset",
        "name": "Study",
        "creator": {"name": "A Person"},
    }
    results = _assess(
        tmp_path, {"study-info.json": json.dumps(document), "trials.csv": _measurement_rows()}
    )
    assert results["F2-METADATA-PRESENT"]["result"] == "pass"
    assert results["F4-METADATA-MACHINE-READABLE"]["result"] == "pass"
    assert results["R1.2-PROVENANCE-DECLARED"]["result"] == "pass"
    assert results["I2-VOCABULARY-REFERENCED"]["result"] == "pass"


def test_a_registry_csv_is_still_read_as_text_and_judged_on_what_it_holds(tmp_path: Path):
    """A CSV's bytes are text, so R1.2 gets a real reading and may fail on it."""
    results = _assess(tmp_path, {"registry.csv": _registry_rows()})
    assert results["F2-METADATA-PRESENT"]["result"] == "pass"
    assert results["R1.2-PROVENANCE-DECLARED"]["result"] == "fail"
    assert "registry.csv" in results["R1.2-PROVENANCE-DECLARED"]["rationale"]


def test_a_worksheet_registry_is_recognised_and_read_as_a_part_of_its_workbook(tmp_path: Path):
    openpyxl = pytest.importorskip("openpyxl")

    root = tmp_path / "wb" / "dataset"
    root.mkdir(parents=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Animals"
    sheet.append(["animal_id", "batch", "sex", "group"])
    for index in range(12):
        sheet.append(
            [
                f"A{index:03d}",
                index % 3 + 1,
                "MF"[index % 2],
                "control" if index % 2 else "treated",
            ]
        )
    book.save(root / "IDs Batch Sex Group.xlsx")

    result = ingest(root, tmp_path / "wb" / "out")
    entry = result.manifest["metadata_files"][0]
    assert entry["path"] == "IDs Batch Sex Group.xlsx#Animals"
    assert entry["file"] == "IDs Batch Sex Group.xlsx"
    assert entry["kind"] == "embedded"
    assert entry["basis"]["sheet"] == "Animals"

    service = DatasetService(result.output_dir, mode="fair-deterministic")
    served = service.get_metadata(entry["path"])
    # The worksheet's bytes are not a file of their own; the service says so
    # rather than serving the whole workbook as if it were the record.
    assert served["content"] is None
    assert "carried inside" in served["content_withheld"]

    results = {item["rule_id"]: item for item in service.run_fair_check()["results"]}
    assert results["F2-METADATA-PRESENT"]["result"] == "pass"
    for rule_id in ("F3-METADATA-LINKS-DATA", "I2-VOCABULARY-REFERENCED"):
        assert results[rule_id]["result"] == "unknown"
        assert "could not be read" in results[rule_id]["rationale"]
    assert results["R1.2-PROVENANCE-DECLARED"]["result"] == "unknown"


def test_a_dataset_whose_only_candidate_was_never_opened_reads_as_unknown(
    tmp_path: Path, monkeypatch
):
    """'No metadata' about a file nobody could open is a claim, not an observation."""
    openpyxl = pytest.importorskip("openpyxl")
    from data2agent.readers import workbook as workbook_reader

    root = tmp_path / "cand" / "dataset"
    root.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.active.append(["animal_id", "group"])
    book.active.append(["A001", "control"])
    book.save(root / "registry.xlsx")

    monkeypatch.setattr(workbook_reader, "available", lambda: False)
    result = ingest(root, tmp_path / "cand" / "out")

    service = DatasetService(result.output_dir, mode="fair-deterministic")
    results = {item["rule_id"]: item for item in service.run_fair_check()["results"]}
    presence = results["F2-METADATA-PRESENT"]
    assert presence["result"] == "unknown"
    assert "could not be examined" in presence["rationale"]
    assert results["R1.2-PROVENANCE-DECLARED"]["result"] == "unknown"


def test_the_example_dataset_is_unchanged_by_content_recognition(ingested):
    """The filename rule still settles everything it settled before."""
    service = DatasetService(ingested.output_dir, mode="fair-deterministic")
    results = {item["rule_id"]: item for item in service.run_fair_check()["results"]}
    assert results["F2-METADATA-PRESENT"]["result"] == "pass"
    assert {item["path"] for item in ingested.manifest["metadata_files"]} == {
        "README.md",
        "dataset_description.json",
    }
    assert all(
        item["recognised_by"] == "filename_convention"
        for item in ingested.manifest["metadata_files"]
    )


def test_one_readable_file_does_not_resolve_another_that_was_never_opened(
    tmp_path: Path, monkeypatch
):
    """Recognising a README says nothing about a workbook nobody could open.

    The uncertainty branches keyed on `unread_metadata()` alone, which holds
    only entries that were RECOGNISED and then could not be decoded. A file no
    recogniser managed to examine is a different population, and once any
    recognised metadata was readable the branch was skipped and the checks
    returned `fail` -- asserting an absence across bytes nobody read.
    """
    openpyxl = pytest.importorskip("openpyxl")
    from data2agent.readers import workbook as workbook_reader

    root = tmp_path / "mixed" / "dataset"
    root.mkdir(parents=True)
    # Readable, filename-recognised, and deliberately silent on every question
    # the four checks ask: no data file named, no vocabulary, no licence, no
    # author. Each check therefore reaches its verdict branch.
    (root / "README.md").write_text("# Study\n\nBehavioural sessions.\n", encoding="utf-8")
    (root / "observations.csv").write_text(_measurement_rows(), encoding="utf-8")
    book = openpyxl.Workbook()
    book.active.append(["animal_id", "group"])
    book.active.append(["A001", "control"])
    book.save(root / "registry.xlsx")

    monkeypatch.setattr(workbook_reader, "available", lambda: False)
    result = ingest(root, tmp_path / "mixed" / "out")

    assert [item["path"] for item in result.manifest["metadata_files"]] == ["README.md"]
    candidates = [item["path"] for item in result.manifest["metadata_candidates"]]
    assert candidates == ["registry.xlsx"], "the workbook must be a candidate, not a recognition"

    service = DatasetService(result.output_dir, mode="fair-deterministic")
    results = {item["rule_id"]: item for item in service.run_fair_check()["results"]}
    for rule_id in (
        "F3-METADATA-LINKS-DATA",
        "I2-VOCABULARY-REFERENCED",
        "R1.1-LICENCE-DECLARED",
        "R1.2-PROVENANCE-DECLARED",
    ):
        assert results[rule_id]["result"] == "unknown", rule_id
        assert "never examined" in results[rule_id]["rationale"], rule_id
        # The candidate is named in the evidence, not merely in the prose, so an
        # agent can act on which file left the verdict open.
        cited = [
            entry
            for entry in results[rule_id]["evidence"]
            if isinstance(entry, dict) and entry["check"] == "metadata.candidate"
        ]
        assert cited, rule_id
        assert [item["path"] for item in cited[0]["result"]] == ["registry.xlsx"], rule_id
        # Nothing was recognised-then-unread here, and no empty list is cited as
        # though it described something.
        assert not [
            entry
            for entry in results[rule_id]["evidence"]
            if isinstance(entry, dict) and entry["check"] == "metadata.content-signature"
        ], rule_id


def test_a_candidate_only_rationale_reads_as_a_sentence(tmp_path: Path, monkeypatch):
    """The detail string joined its two clauses with a leading 'and'."""
    pytest.importorskip("openpyxl")
    import openpyxl

    from data2agent.readers import workbook as workbook_reader

    root = tmp_path / "prose" / "dataset"
    root.mkdir(parents=True)
    (root / "README.md").write_text("# Study\n", encoding="utf-8")
    (root / "observations.csv").write_text(_measurement_rows(), encoding="utf-8")
    book = openpyxl.Workbook()
    book.active.append(["animal_id", "group"])
    book.save(root / "registry.xlsx")

    monkeypatch.setattr(workbook_reader, "available", lambda: False)
    result = ingest(root, tmp_path / "prose" / "out")
    service = DatasetService(result.output_dir, mode="fair-deterministic")
    results = {item["rule_id"]: item for item in service.run_fair_check()["results"]}

    rationale = results["R1.2-PROVENANCE-DECLARED"]["rationale"]
    assert "but and" not in rationale and "; and" not in rationale
    assert "1 file(s) were never examined" in rationale


def test_the_standard_is_cited_by_what_established_it(tmp_path: Path):
    """R1.3 passed on a content recognition while citing only a README.

    `community_standard` cited every `metadata.file-convention` claim, so a BIDS
    document recognised from its content contributed the verdict while the
    evidence pointed at a README -- documentation, which is explicitly not a
    community standard.
    """
    root = tmp_path / "std" / "dataset"
    root.mkdir(parents=True)
    (root / "README.md").write_text("# Study\n", encoding="utf-8")
    # Not named dataset_description.json, so only the content rule can reach it.
    (root / "study_meta.json").write_text(
        json.dumps({"Name": "APA", "BIDSVersion": "1.8.0"}), encoding="utf-8"
    )
    (root / "observations.csv").write_text(_measurement_rows(), encoding="utf-8")

    result = ingest(root, tmp_path / "std" / "out")
    carrier = next(
        item for item in result.manifest["metadata_files"] if item["convention"] == "bids"
    )
    assert carrier["path"] == "study_meta.json"
    assert carrier["recognised_by"] == "content"

    service = DatasetService(result.output_dir, mode="fair-deterministic")
    results = {item["rule_id"]: item for item in service.run_fair_check()["results"]}
    standard = results["R1.3-COMMUNITY-STANDARD"]
    assert standard["result"] == "pass"

    evidence = json.loads((result.output_dir / "evidence.json").read_text(encoding="utf-8"))
    by_id = {claim["claim_id"]: claim for claim in evidence["claims"]}
    cited = [by_id[claim_id] for claim_id in standard["evidence"] if claim_id in by_id]
    assert cited, "the verdict must cite ledger claims, not restate the standard"
    assert {item["subject"] for item in cited} == {"study_meta.json"}
    assert any(
        entry["check"] == "metadata.content-signature"
        for claim in cited
        for entry in claim["evidence"]
    )
    # The README is recognised metadata, but it did not establish the standard.
    assert "README.md" not in {item["subject"] for item in cited}
